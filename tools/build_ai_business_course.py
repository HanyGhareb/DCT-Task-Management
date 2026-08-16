from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR


OUT = Path("deliverables/ai-for-business-course")
OUT.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(17, 31, 52)
INK = RGBColor(31, 45, 61)
TEAL = RGBColor(18, 151, 147)
CYAN = RGBColor(56, 189, 208)
GOLD = RGBColor(242, 177, 52)
CORAL = RGBColor(239, 106, 83)
WHITE = RGBColor(255, 255, 255)
MIST = RGBColor(242, 247, 249)
PALE = RGBColor(225, 241, 242)
GRAY = RGBColor(100, 116, 139)
LIGHT_GRAY = RGBColor(220, 228, 234)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "AI for Everyday Business"
prs.core_properties.subject = "Practical AI productivity course for non-technical business professionals"
prs.core_properties.author = "AI for Everyday Business"


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def textbox(slide, text, x, y, w, h, size=20, color=INK, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return box


def add_runs(box, parts, size=20, align=PP_ALIGN.LEFT):
    tf = box.text_frame; tf.clear()
    p = tf.paragraphs[0]; p.alignment = align
    for text, color, bold in parts:
        r = p.add_run(); r.text = text
        r.font.name = "Aptos"; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold


def base_slide(section="AI FOR EVERYDAY BUSINESS", num=None, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = NAVY if dark else WHITE
    if not dark:
        rect(slide, 0, 0, .13, 7.5, TEAL)
        textbox(slide, section.upper(), .55, .22, 7.5, .25, 9, TEAL, True)
        if num is not None:
            textbox(slide, f"{num:02d}", 12.15, 7.05, .55, .22, 9, GRAY, True, align=PP_ALIGN.RIGHT)
    return slide


def title(slide, text, subtitle=None, dark=False):
    color = WHITE if dark else NAVY
    textbox(slide, text, .55, .62, 12.0, .72, 28, color, True)
    if subtitle:
        textbox(slide, subtitle, .58, 1.35, 11.8, .48, 13, CYAN if dark else GRAY)


def bullet_list(slide, items, x, y, w, h, size=18, color=INK, accent=TEAL, gap=7):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.05); tf.margin_right = Inches(.03)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.level = 0
        p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = color
        p.space_after = Pt(gap); p.line_spacing = 1.05
        p._p.get_or_add_pPr().insert(0, p._p._new_buChar()) if False else None
        p.text = "●  " + item
        p.runs[0].font.color.rgb = color
    return box


def card(slide, x, y, w, h, heading, body, accent=TEAL, number=None):
    rect(slide, x, y, w, h, MIST, True, LIGHT_GRAY)
    rect(slide, x, y, .09, h, accent, True, accent)
    if number:
        rect(slide, x+.25, y+.25, .5, .5, accent, True)
        textbox(slide, str(number), x+.25, y+.25, .5, .5, 15, WHITE, True,
                align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        hx = x+.9
    else:
        hx = x+.28
    textbox(slide, heading, hx, y+.22, w-(hx-x)-.2, .34, 16, NAVY, True)
    textbox(slide, body, x+.28, y+.72, w-.5, h-.87, 11.5, GRAY)


def footer(slide, text="Practical • Responsible • Human-led"):
    textbox(slide, text, .58, 7.05, 6.5, .2, 8.5, GRAY)


# 1 Cover
s = base_slide(dark=True)
rect(s, 0, 0, .18, 7.5, TEAL)
rect(s, 9.45, -.5, 4.5, 8.5, RGBColor(22, 52, 74), True)
for x, y, c, sz in [(10.1,1.0,TEAL,1.1),(11.6,2.1,GOLD,.65),(9.9,3.4,CYAN,.55),(11.2,4.6,CORAL,.9),(10.0,5.8,TEAL,.55)]:
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(sz), Inches(sz))
    shp.fill.solid(); shp.fill.fore_color.rgb=c; shp.line.color.rgb=c
textbox(s, "AI FOR", .72, 1.15, 7.8, .5, 18, CYAN, True)
textbox(s, "Everyday Business", .68, 1.72, 8.5, 1.08, 38, WHITE, True)
textbox(s, "Practical productivity for leaders and professionals", .72, 3.02, 7.6, .55, 19, WHITE)
textbox(s, "Emails  •  Documents  •  Data  •  Presentations  •  Planning  •  Meetings", .72, 4.05, 8.0, .8, 14, RGBColor(198,218,226), True)
rect(s, .72, 5.4, 3.15, .55, TEAL, True)
textbox(s, "NON-TECHNICAL COURSE", .83, 5.54, 2.92, .24, 10, WHITE, True, align=PP_ALIGN.CENTER)
textbox(s, "Course deck • 12–15 hours", .75, 6.55, 5, .3, 11, RGBColor(166,191,203))

# 2 Course promise
s=base_slide("COURSE OVERVIEW",2); title(s,"The course promise","Use AI to improve work—not to surrender judgment.")
card(s,.6,2.1,3.85,3.45,"SAVE TIME","Draft, summarize, classify and organize routine work faster.",TEAL,"01")
card(s,4.75,2.1,3.85,3.45,"IMPROVE QUALITY","Communicate clearly, analyze information and develop stronger first drafts.",CYAN,"02")
card(s,8.9,2.1,3.85,3.45,"WORK RESPONSIBLY","Protect information, verify important outputs and keep people accountable.",GOLD,"03")
textbox(s,"Designed for executives, managers, team leaders and employees—no coding required.",.8,6.25,11.6,.5,16,NAVY,True,align=PP_ALIGN.CENTER); footer(s)

# 3 outcomes
s=base_slide("COURSE OVERVIEW",3); title(s,"What participants will be able to do")
left=["Choose the right AI tool for a business task","Write clear, reusable prompts","Handle email and documents efficiently","Analyze spreadsheets in natural language"]
right=["Create stronger presentations","Build weekly plans and meeting workflows","Research and verify information","Use AI safely within company policy"]
bullet_list(s,left,.75,1.85,5.7,4.7,18); bullet_list(s,right,6.8,1.85,5.7,4.7,18,accent=CYAN)
rect(s,.7,6.25,12,.52,PALE,True); textbox(s,"Outcome: a repeatable AI-assisted workflow for each participant’s real job.",.9,6.36,11.6,.28,14,TEAL,True,align=PP_ALIGN.CENTER); footer(s)

# 4 format
s=base_slide("COURSE OVERVIEW",4); title(s,"Recommended learning format","12–15 hours • flexible delivery")
for i,(h,b,c) in enumerate([("20%","EXPLAIN\nConcepts in plain language",TEAL),("30%","DEMONSTRATE\nReal business workflows",CYAN),("40%","PRACTISE\nGuided hands-on tasks",GOLD),("10%","REVIEW\nKnowledge checks",CORAL)]):
    x=.65+i*3.15; rect(s,x,2.0,2.75,3.2,MIST,True,LIGHT_GRAY); rect(s,x,2.0,2.75,.16,c,True,c)
    textbox(s,h,x+.2,2.48,2.35,.62,30,c,True,align=PP_ALIGN.CENTER)
    textbox(s,b,x+.25,3.35,2.25,1.0,13,NAVY,True,align=PP_ALIGN.CENTER)
textbox(s,"Delivery options: 6 weekly sessions  •  10 self-paced modules  •  2-day workshop  •  4-hour executive edition",.75,5.85,11.8,.6,14,GRAY,align=PP_ALIGN.CENTER); footer(s)

# 5 roadmap
s=base_slide("COURSE ROADMAP",5); title(s,"Ten modules—from curiosity to capability")
mods=[("01","AI at work"),("02","Prompting"),("03","Email"),("04","Documents"),("05","Data & Excel"),("06","Presentations"),("07","Weekly planning"),("08","Meetings"),("09","Business roles"),("10","Responsible AI")]
for i,(n,lbl) in enumerate(mods):
    col=i%5; row=i//5; x=.65+col*2.52; y=1.8+row*2.25
    rect(s,x,y,2.2,1.62,MIST,True,LIGHT_GRAY); textbox(s,n,x+.18,y+.16,.5,.3,12,TEAL,True)
    textbox(s,lbl,x+.18,y+.64,1.84,.6,16,NAVY,True)
textbox(s,"CAPSTONE",5.3,6.25,1.4,.28,11,WHITE,True,align=PP_ALIGN.CENTER); rect(s,5.1,6.12,1.8,.55,TEAL,True)
textbox(s,"CAPSTONE",5.3,6.26,1.4,.2,10,WHITE,True,align=PP_ALIGN.CENTER); textbox(s,"Redesign one real workflow",7.05,6.2,3.5,.3,14,NAVY,True); footer(s)

# 6 AI basics
s=base_slide("MODULE 1 • AI AT WORK",6); title(s,"Know where AI helps—and where it fails")
card(s,.65,1.75,5.8,4.7,"AI IS STRONG AT","Drafting and rewriting\nSummarizing and classifying\nComparing options\nBrainstorming and structuring\nAnalyzing structured data",TEAL)
card(s,6.85,1.75,5.8,4.7,"AI NEEDS HUMAN CONTROL","Factual accuracy\nMissing business context\nHigh-stakes decisions\nCurrent or authoritative information\nEthical and legal accountability",CORAL)
footer(s)

# 7 prompt framework
s=base_slide("MODULE 2 • PROMPTING",7); title(s,"CRAFT a better business prompt","Five ingredients turn a vague request into usable work.")
craft=[("C","CONTEXT","What is the situation?",TEAL),("R","ROLE","What expertise is useful?",CYAN),("A","ACTION","What must AI do?",GOLD),("F","FORMAT","How should it respond?",CORAL),("T","TESTS","What must be true?",TEAL)]
for i,(letter,h,b,c) in enumerate(craft):
    x=.5+i*2.56; rect(s,x,2.05,2.25,3.4,MIST,True,LIGHT_GRAY)
    shp=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+.72),Inches(2.35),Inches(.8),Inches(.8)); shp.fill.solid();shp.fill.fore_color.rgb=c;shp.line.color.rgb=c
    textbox(s,letter,x+.72,2.35,.8,.8,25,WHITE,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    textbox(s,h,x+.18,3.48,1.9,.3,13,NAVY,True,align=PP_ALIGN.CENTER)
    textbox(s,b,x+.22,4.05,1.82,.65,11.5,GRAY,align=PP_ALIGN.CENTER)
rect(s,.75,5.95,11.8,.62,PALE,True); textbox(s,"Tip: ask AI to identify missing information before it writes the final answer.",1.0,6.1,11.3,.25,14,TEAL,True,align=PP_ALIGN.CENTER); footer(s)

# 8 prompt example
s=base_slide("MODULE 2 • PROMPTING",8); title(s,"From vague request to reliable instruction")
textbox(s,"WEAK",.75,1.72,1,.3,11,CORAL,True); rect(s,.72,2.08,4.0,1.25,RGBColor(253,238,235),True,RGBColor(248,200,191)); textbox(s,"“Summarize this meeting.”",1.0,2.47,3.5,.35,18,NAVY,True,align=PP_ALIGN.CENTER)
textbox(s,"CRAFTED",5.15,1.72,1.2,.3,11,TEAL,True); rect(s,5.1,2.08,7.55,3.52,PALE,True,RGBColor(180,220,220))
textbox(s,"You are an executive assistant. Using only the meeting notes provided, prepare a concise management summary. Include decisions, action items, owners, deadlines and unresolved issues. Present actions in a table. Clearly mark missing owners or dates; do not invent them.",5.45,2.42,6.86,2.65,16,INK)
textbox(s,"PRACTICE",.75,4.25,1.2,.3,11,GOLD,True); textbox(s,"Rewrite three weak prompts using CRAFT.",.78,4.7,3.7,1.0,18,NAVY,True); footer(s)

# 9 email
s=base_slide("MODULE 3 • EMAIL",9); title(s,"Turn the inbox into decisions and actions")
items=[("SUMMARIZE","Long threads → five key points"),("DRAFT","New messages and replies"),("ADAPT","Formal, friendly, firm or diplomatic"),("EXTRACT","Decisions, owners and deadlines"),("FOLLOW UP","Reminders and meeting invitations"),("CHECK","Facts, recipients, tone and attachments")]
for i,(h,b) in enumerate(items):
    x=.65+(i%3)*4.18; y=1.75+(i//3)*2.22; card(s,x,y,3.75,1.82,h,b,[TEAL,CYAN,GOLD,CORAL,TEAL,CYAN][i])
footer(s)

# 10 email lab
s=base_slide("MODULE 3 • EMAIL",10); title(s,"Lab: tame a 15-message project thread")
steps=["Produce a five-line summary","Extract decisions and commitments","Create an owner / action / deadline table","Draft a concise reply to all","Prepare a calendar invitation"]
for i,st in enumerate(steps):
    y=1.65+i*.9; rect(s,.78,y,.58,.58,TEAL if i<4 else GOLD,True); textbox(s,str(i+1),.78,y,.58,.58,17,WHITE,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE); textbox(s,st,1.62,y+.1,6.4,.35,17,NAVY,True)
rect(s,8.5,1.72,3.85,4.55,NAVY,True); textbox(s,"BEFORE SEND",8.9,2.1,3.05,.35,14,CYAN,True,align=PP_ALIGN.CENTER)
bullet_list(s,["Recipients","Names and facts","Dates and promises","Tone","Attachments","Confidential data"],9.0,2.75,2.9,2.9,14,WHITE,gap=8)
footer(s)

# 11 docs
s=base_slide("MODULE 4 • DOCUMENTS",11); title(s,"Move from rough notes to polished documents")
flow=[("NOTES","Raw ideas\nEmails\nTranscripts"),("STRUCTURE","Outline\nAudience\nPurpose"),("DRAFT","Report\nProposal\nPolicy"),("REFINE","Tone\nClarity\nConsistency"),("VERIFY","Facts\nObligations\nRisks")]
for i,(h,b) in enumerate(flow):
    x=.42+i*2.55; rect(s,x,2.05,2.13,2.75,MIST,True,LIGHT_GRAY); textbox(s,h,x+.15,2.42,1.83,.3,13,TEAL,True,align=PP_ALIGN.CENTER); textbox(s,b,x+.25,3.15,1.63,1.05,14,NAVY,align=PP_ALIGN.CENTER)
    if i<4: textbox(s,"›",x+2.16,2.95,.35,.5,28,GOLD,True,align=PP_ALIGN.CENTER)
textbox(s,"Exercise: one set of notes → executive brief • customer update • staff announcement • action list",.8,5.65,11.7,.62,15,NAVY,True,align=PP_ALIGN.CENTER); footer(s)

# 12 Excel
s=base_slide("MODULE 5 • DATA & EXCEL",12); title(s,"Ask business questions—not just formula questions")
questions=["Is the data complete and consistent?","What changed—and why might it matter?","Where are the exceptions or unusual values?","How do actual results compare with target?","Which chart communicates the finding best?","Can I reproduce the key calculation?"]
for i,q in enumerate(questions):
    x=.68+(i%2)*6.15; y=1.65+(i//2)*1.45
    rect(s,x,y,5.72,1.08,MIST,True,LIGHT_GRAY); rect(s,x,y,.13,1.08,[TEAL,CYAN,GOLD][i//2],True); textbox(s,q,x+.35,y+.27,5.05,.52,15.5,NAVY,True)
rect(s,.78,6.18,11.75,.5,RGBColor(253,238,235),True); textbox(s,"Guardrail: AI can suggest an analysis; the business owner must verify material figures.",1.0,6.3,11.3,.25,13,CORAL,True,align=PP_ALIGN.CENTER); footer(s)

# 13 data lab
s=base_slide("MODULE 5 • DATA & EXCEL",13); title(s,"Lab: turn sales data into management insight")
labs=[("1","CHECK","Missing values, duplicates, types"),("2","COMPARE","Month, region and category"),("3","VISUALIZE","Select and create a clear chart"),("4","EXPLAIN","Draft five management observations"),("5","VERIFY","Recalculate three critical figures")]
for i,(n,h,b) in enumerate(labs):
    x=.46+i*2.55; card(s,x,2.0,2.25,3.55,h,b,[TEAL,CYAN,GOLD,CORAL,TEAL][i],n)
textbox(s,"Good analysis distinguishes evidence from interpretation—and correlation from causation.",.85,6.15,11.5,.45,15,NAVY,True,align=PP_ALIGN.CENTER); footer(s)

# 14 presentations
s=base_slide("MODULE 6 • PRESENTATIONS",14); title(s,"Build a decision story—not a pile of slides")
story=["SITUATION","EVIDENCE","FINDINGS","OPTIONS","RECOMMENDATION","PLAN","DECISION"]
for i,label in enumerate(story):
    x=.42+i*1.82; c=TEAL if i in (0,6) else (CYAN if i<4 else GOLD)
    rect(s,x,2.28,1.52,1.05,c,True); textbox(s,label,x+.08,2.62,1.36,.24,10.5,WHITE,True,align=PP_ALIGN.CENTER)
    if i<6: textbox(s,"›",x+1.52,2.52,.3,.4,22,GRAY,True,align=PP_ALIGN.CENTER)
bullet_list(s,["Conclusion-led slide titles","One message per slide","Evidence for every important claim","Speaker notes for delivery","Brand and accessibility checks"],1.1,4.1,11.2,1.9,16)
footer(s)

# 15 planning
s=base_slide("MODULE 7 • WEEKLY PLANNING",15); title(s,"Convert an overloaded list into a realistic week")
for i,(h,b,c) in enumerate([("PRIORITIZE","Choose three weekly outcomes",TEAL),("BLOCK","Reserve focus and preparation time",CYAN),("BALANCE","Keep 15% contingency",GOLD),("DECIDE","Do • delegate • defer • decline",CORAL)]):
    x=.65+i*3.12; card(s,x,2.0,2.76,3.7,h,b,c,str(i+1))
textbox(s,"AI proposes the schedule. You decide what deserves your time.",1.0,6.1,11.2,.45,16,NAVY,True,align=PP_ALIGN.CENTER); footer(s)

# 16 planning prompt
s=base_slide("MODULE 7 • WEEKLY PLANNING",16); title(s,"Reusable prompt: weekly planning coach")
rect(s,.75,1.62,11.8,3.65,PALE,True,RGBColor(180,220,220))
textbox(s,"Act as my business productivity coach. Using my goals, deadlines, meetings and available hours, create a realistic weekly plan. Separate fixed commitments from flexible tasks. Include three priorities, daily time blocks, preparation and follow-up time, and 15% contingency. Identify conflicts and tasks to delegate. Ask questions where important information is missing.",1.15,2.02,11.0,2.75,18,INK)
textbox(s,"INPUTS TO PROVIDE",.8,5.75,1.8,.3,11,TEAL,True)
textbox(s,"Goals  •  deadlines  •  meetings  •  working hours  •  energy constraints  •  delegation options",2.65,5.72,9.5,.5,14,NAVY,True); footer(s)

# 17 meetings
s=base_slide("MODULE 8 • MEETINGS",17); title(s,"Use AI across the meeting lifecycle")
stages=[("NEED?","Could an email solve it?"),("PREPARE","Objective, people, agenda"),("RUN","Questions and notes"),("CLOSE","Decisions and owners"),("FOLLOW","Actions and reminders")]
for i,(h,b) in enumerate(stages):
    x=.5+i*2.55; c=[TEAL,CYAN,GOLD,CORAL,TEAL][i]
    shp=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+.61),Inches(2.0),Inches(1.0),Inches(1.0));shp.fill.solid();shp.fill.fore_color.rgb=c;shp.line.color.rgb=c
    textbox(s,str(i+1),x+.61,2.0,1,1,24,WHITE,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    textbox(s,h,x+.15,3.35,1.9,.3,13,NAVY,True,align=PP_ALIGN.CENTER); textbox(s,b,x+.05,3.85,2.1,.75,12,GRAY,align=PP_ALIGN.CENTER)
    if i<4: textbox(s,"→",x+2.05,2.25,.45,.4,20,LIGHT_GRAY,True,align=PP_ALIGN.CENTER)
rect(s,.78,5.48,11.78,.7,RGBColor(253,238,235),True); textbox(s,"Recording or transcribing? Obtain consent and follow organizational privacy rules.",1.0,5.68,11.3,.28,14,CORAL,True,align=PP_ALIGN.CENTER); footer(s)

# 18 roles
s=base_slide("MODULE 9 • BUSINESS ROLES",18); title(s,"Apply AI where work actually happens")
roles=[("LEADERS","Briefings • options • communication",TEAL),("HR","Training • surveys • job drafts",CYAN),("SALES","Preparation • proposals • feedback",GOLD),("FINANCE","Variance • controls • scenarios",CORAL),("OPERATIONS","Plans • risks • procedures",TEAL)]
for i,(h,b,c) in enumerate(roles):
    x=.45+i*2.55; card(s,x,2.05,2.25,3.45,h,b,c)
textbox(s,"High-impact employment, financial, legal and customer decisions remain human decisions.",.85,6.02,11.6,.5,14,NAVY,True,align=PP_ALIGN.CENTER); footer(s)

# 19 responsible AI
s=base_slide("MODULE 10 • RESPONSIBLE AI",19); title(s,"A simple rule before every prompt")
rect(s,.8,1.65,11.7,1.05,NAVY,True); textbox(s,"Is both the information and the AI system approved by my organization?",1.15,1.97,11,.38,19,WHITE,True,align=PP_ALIGN.CENTER)
guards=[("PROTECT","Personal, customer, employee and confidential data",TEAL),("VERIFY","Facts, sources, calculations and quotations",CYAN),("CONTROL","Human approval for material actions",GOLD),("DOCUMENT","Inputs, outputs and decisions when required",CORAL)]
for i,(h,b,c) in enumerate(guards):
    x=.68+i*3.1; card(s,x,3.35,2.72,2.6,h,b,c)
footer(s)

# 20 verification
s=base_slide("MODULE 10 • RESPONSIBLE AI",20); title(s,"The nine-point verification check")
checks=["Names, dates and figures","Source-bound—not invented","Important omissions","Audience and tone","Confidential information","Company policy","Bias or unfair impact","Expert approval needed","I accept responsibility"]
for i,c in enumerate(checks):
    col=i%3; row=i//3; x=.7+col*4.12; y=1.65+row*1.48
    rect(s,x,y,3.75,1.1,MIST,True,LIGHT_GRAY); rect(s,x+.22,y+.27,.48,.48,TEAL,True); textbox(s,"✓",x+.22,y+.27,.48,.48,16,WHITE,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE); textbox(s,c,x+.88,y+.28,2.62,.5,14,NAVY,True)
footer(s)

# 21 capstone
s=base_slide("CAPSTONE PROJECT",21); title(s,"Redesign one real business workflow")
deliver=[("CURRENT","Process, pain points and time"),("DESIGN","AI-assisted workflow and prompts"),("CONTROL","Human reviews and data risks"),("PROVE","Example outputs and expected gain"),("MEASURE","30-day adoption and quality plan")]
for i,(h,b) in enumerate(deliver):
    x=.45+i*2.55; card(s,x,2.0,2.25,3.65,h,b,[TEAL,CYAN,GOLD,CORAL,TEAL][i],str(i+1))
textbox(s,"Examples: management reporting • customer inquiries • meeting follow-up • onboarding • project status",.7,6.13,12,.38,14,GRAY,align=PP_ALIGN.CENTER); footer(s)

# 22 assessment
s=base_slide("ASSESSMENT",22); title(s,"Assess judgment—not prompt memorization")
assess=[("Knowledge checks",15),("Email",10),("Documents",10),("Spreadsheet",15),("Presentation",10),("Weekly planning",10),("Responsible AI",10),("Capstone",20)]
maxw=6.0
for i,(label,val) in enumerate(assess):
    col=i//4; row=i%4; x=.75+col*6.25; y=1.62+row*1.15
    textbox(s,label,x,y,2.15,.3,13,NAVY,True); rect(s,x+2.15,y+.02,2.95,.28,LIGHT_GRAY,True); rect(s,x+2.15,y+.02,2.95*val/20,.28,TEAL if col==0 else CYAN,True); textbox(s,f"{val}%",x+5.25,y, .55,.3,12,GRAY,True,align=PP_ALIGN.RIGHT)
rect(s,.8,6.35,11.7,.48,PALE,True); textbox(s,"Rubric: usefulness • accuracy • verification • communication • safety",1.0,6.46,11.3,.24,13,TEAL,True,align=PP_ALIGN.CENTER); footer(s)

# 23 resources
s=base_slide("COURSE RESOURCES",23); title(s,"Give participants tools they can reuse Monday morning")
resources=["Participant workbook","CRAFT prompt card","50–100 prompt library","Email prompt pack","Excel analysis pack","Document & slide pack","Weekly plan template","Meeting templates","Verification checklist","Data decision guide","Responsible AI policy","Fictional practice files"]
for i,r in enumerate(resources):
    col=i%3; row=i//3; x=.65+col*4.18; y=1.55+row*1.22
    rect(s,x,y,3.75,.9,MIST,True,LIGHT_GRAY); textbox(s,f"{i+1:02d}",x+.18,y+.25,.4,.3,11,TEAL,True); textbox(s,r,x+.75,y+.22,2.75,.45,13,NAVY,True)
textbox(s,"One fictional company can connect every exercise into a coherent business story.",.85,6.62,11.6,.3,13,GRAY,align=PP_ALIGN.CENTER); footer(s)

# 24 rollout
s=base_slide("IMPLEMENTATION",24); title(s,"Make the learning stick after the course")
for i,(n,h,b,c) in enumerate([("01","BASELINE","Measure speed, quality and confidence before training",TEAL),("02","PRACTISE","Use fictional, realistic files and show AI failures",CYAN),("03","ADOPT","Apply one workflow in the participant’s real role",GOLD),("04","REVIEW","Return after 30 days to measure impact and risk",CORAL)]):
    x=.65+i*3.12; card(s,x,2.0,2.76,3.7,h,b,c,n)
textbox(s,"Keep the workflow lessons stable; update short tool demonstrations as products change.",.75,6.18,11.85,.38,14,NAVY,True,align=PP_ALIGN.CENTER); footer(s)

# 25 sources
s=base_slide("REFERENCES",25); title(s,"Official product references","Feature availability varies by plan, account, region, language and organizational settings.")
refs=[
    ("Microsoft Copilot overview","support.microsoft.com/en-us/office/what-you-can-do-with-microsoft-copilot"),
    ("Copilot in Excel","support.microsoft.com/en-us/excel/copilot/get-started-with-copilot-in-excel"),
    ("Gemini in Gmail","support.google.com/mail/answer/14355636"),
    ("Gemini in Docs, Sheets & Slides","support.google.com/docs/answer/15123226"),
    ("Google Workspace AI privacy","support.google.com/docs/answer/14615114"),
    ("Yahoo Mail AI features","help.yahoo.com/kb/SLN36900.html"),
    ("ChatGPT data analysis","help.openai.com/en/articles/9213685"),
    ("ChatGPT capabilities overview","help.openai.com/en/articles/9260256"),
]
for i,(h,u) in enumerate(refs):
    col=i//4; row=i%4; x=.72+col*6.2; y=1.62+row*1.18
    textbox(s,h,x,y,5.4,.28,13,NAVY,True); textbox(s,u,x,y+.37,5.5,.3,9.5,TEAL)
textbox(s,"References checked August 2026 • Always confirm current product documentation before recording demonstrations.",.85,6.52,11.6,.35,11,GRAY,align=PP_ALIGN.CENTER); footer(s)

# 26 close
s=base_slide(dark=True)
rect(s,0,0,.18,7.5,TEAL); textbox(s,"THE GOAL",.75,1.05,3,.35,14,CYAN,True)
textbox(s,"Better work.\nBetter judgment.\nMore time for people.",.72,1.65,9.4,2.75,36,WHITE,True)
rect(s,.75,5.28,5.4,.08,TEAL)
textbox(s,"AI for Everyday Business",.75,5.65,6,.4,18,WHITE,True)
textbox(s,"Practical • Responsible • Human-led",.75,6.23,6,.3,12,RGBColor(174,198,209))

path=OUT / "AI_for_Everyday_Business_Course.pptx"
prs.save(path)
print(path)
