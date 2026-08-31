#!/usr/bin/env python3
"""UAT runner — Finance KPIs (App 213). Playwright + ORDS hybrid (model: uat_run_cc.py).

Executes the KPI UAT cases end-to-end (browser UI + API), captures one evidence
screenshot per case, and emits the Admin-convention package:

  final apps/KPI/UAT/UAT_KPI_TestScript.xlsx                      (reusable master)
  final apps/KPI/UAT/UAT_KPI_round<N>-dd-mm-yyyy/
      UAT_KPI_<dd-Mon-yyyy>-NN.xlsx                               (statuses filled)
      UAT_KPI_Results_<dd-Mon-yyyy>-NN.docx
      evidence_<dd-Mon-yyyy>-NN/case_XX.png

Env: KPI_TOK_ADMIN / KPI_TOK_USER / KPI_TOK_SYS (live sessions), KPI_PORT (dev-proxy),
     KPI_TEST_YEAR (fresh synthetic year), KPI_ROUND (default 1).
"""
import json, os, sys, datetime, urllib.request, urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from playwright.sync_api import sync_playwright

BASE = 'https://gd5cec2eaeb21e3-prod.adb.me-abudhabi-1.oraclecloudapps.com/ords/admin'
APP  = 'http://localhost:' + os.environ.get('KPI_PORT', '8098')
TOKA = os.environ.get('KPI_TOK_ADMIN'); TOKU = os.environ.get('KPI_TOK_USER'); TOKS = os.environ.get('KPI_TOK_SYS')
YEAR = int(os.environ.get('KPI_TEST_YEAR', '2094'))
ROUND = int(os.environ.get('KPI_ROUND', '1'))
if not (TOKA and TOKU and TOKS):
    sys.exit('Set KPI_TOK_ADMIN / KPI_TOK_USER / KPI_TOK_SYS')

TODAY = datetime.date.today()
DMY   = TODAY.strftime('%d-%m-%Y')
DMON  = TODAY.strftime('%d-%b-%Y')
ROOT  = os.path.join(os.path.dirname(__file__), '..', '..', '..', 'final apps', 'KPI', 'UAT')
RDIR  = os.path.join(ROOT, 'UAT_KPI_round%d-%s' % (ROUND, DMY))
EDIR  = os.path.join(RDIR, 'evidence_%s-01' % DMON)
os.makedirs(EDIR, exist_ok=True)

def api(method, path, token=None, body=None, raw=None, mime=None):
    headers = {}
    data = None
    if token: headers['Authorization'] = 'Bearer ' + token
    if body is not None:
        data = json.dumps(body).encode(); headers['Content-Type'] = 'application/json'
    if raw is not None:
        data = raw; headers['Content-Type'] = mime or 'application/octet-stream'
    req = urllib.request.Request(BASE + path, data=data, headers=headers, method=method)
    try:
        r = urllib.request.urlopen(req, timeout=120); code, payload = r.status, r.read()
    except urllib.error.HTTPError as e:
        code, payload = e.code, e.read()
    try: parsed = json.loads(payload) if payload else {}
    except Exception: parsed = {}
    return code, parsed

CASES = []   # (id, area, title, steps, expected, status, actual)
def record(cid, area, title, steps, expected, ok, actual):
    CASES.append((cid, area, title, steps, expected, 'PASS' if ok else 'FAIL', actual))
    print(('PASS  ' if ok else 'FAIL  ') + cid + ' ' + title + ('' if ok else ' -- ' + str(actual)))

with sync_playwright() as p:
    browser = p.chromium.launch(headless=True)
    page = browser.new_context(viewport={'width': 1440, 'height': 900}).new_page()
    page.goto(APP + '/KPI/Jet/index.html')
    page.evaluate("s => localStorage.setItem('ifinance_jet_session', JSON.stringify(s))", {
        'sessionId': TOKA, 'userId': 21, 'username': 'HANY', 'displayName': 'Hany',
        'email': 'uat@test', 'roles': ['KPI_ADMIN', 'KPI_USER']})
    page.goto(APP + '/KPI/Jet/index.html')
    page.wait_for_timeout(5000)

    def shot(cid):
        page.screenshot(path=os.path.join(EDIR, 'case_%s.png' % cid), full_page=False)

    # ---- KPI-01 app boot / scorecard --------------------------------------
    ok = page.locator('.kp-overall').count() == 1 and page.locator('.kp-card').count() == 4
    shot('01')
    record('KPI-01', 'Dashboard', 'Scorecard dashboard renders',
           'Open the KPI app with a valid session', '4 KPI cards + overall weighted tile', ok,
           '%d cards' % page.locator('.kp-card').count())

    # ---- KPI-02 registry --------------------------------------------------
    page.evaluate("window._jetApp.navigate('kpiList')"); page.wait_for_timeout(3000)
    n = page.locator('.data-table tbody tr').count(); shot('02')
    record('KPI-02', 'Registry', 'Registry lists the 4 circular KPIs',
           'Administration > KPI Registry', 'REV_GROWTH, OPT_PLAN, CASH_MGMT, FIN_LAW_COMP listed',
           n == 4, '%d rows' % n)

    # ---- KPI-03 definition completeness (API) -----------------------------
    c, d = api('GET', '/kpi/kpis', TOKU)
    kpis = {i['code']: i for i in d.get('items', [])}
    c, d = api('GET', '/kpi/kpis/%d' % kpis['OPT_PLAN']['kpiId'], TOKU)
    ok = (len(d.get('bands', [])) == 5 and len(d.get('criteria', [])) == 4
          and all(len(cr['levels']) == 5 for cr in d['criteria']) and len(d.get('targets', [])) == 3)
    page.locator('.data-table tbody tr').nth(1).click(); page.wait_for_timeout(3500); shot('03')
    record('KPI-03', 'Registry', 'OPT_PLAN carries full circular content',
           'Open OPT_PLAN in the editor', '5 bands, 4 criteria x 5 maturity levels, 3 targets', ok,
           'bands=%d criteria=%d' % (len(d.get('bands', [])), len(d.get('criteria', []))))

    # ---- KPI-04 admin gating ----------------------------------------------
    c1, _ = api('POST', '/kpi/kpis', TOKU, body={'code': 'X', 'nameEn': 'x', 'frequency': 'ANNUAL', 'calcMethod': 'RATIO_A_OVER_B'})
    c2, _ = api('POST', '/kpi/periods/generate', TOKU, body={'year': YEAR})
    shot('04')
    record('KPI-04', 'Security', 'Non-admin blocked from configuration',
           'As a plain KPI_USER call POST kpis and POST periods/generate', 'Both return 403',
           c1 == 403 and c2 == 403, (c1, c2))

    # ---- KPI-05 period generation -----------------------------------------
    c, _ = api('POST', '/kpi/periods/generate', TOKA, body={'year': YEAR})
    c2, d = api('GET', '/kpi/periods?year=%d' % YEAR, TOKA)
    periods = {(i['type'], i.get('quarter')): i['periodId'] for i in d.get('items', [])}
    shot('05')
    record('KPI-05', 'Periods', 'Measurement calendar generation',
           'KPI_ADMIN generates the %d calendar' % YEAR, '1 annual + 4 quarter periods',
           c == 200 and len(d.get('items', [])) == 5, '%d periods' % len(d.get('items', [])))

    # ---- KPI-06 frequency guard -------------------------------------------
    c, d = api('POST', '/kpi/results/init', TOKU,
               body={'kpiId': kpis['CASH_MGMT']['kpiId'], 'periodId': periods[('ANNUAL', None)]})
    record('KPI-06', 'Results', 'Quarterly KPI rejected on an annual period',
           'Start CASH_MGMT on the annual period', 'HTTP 400 frequency mismatch', c == 400, c)
    shot('06')

    # ---- KPI-07 ratio figures + live scoring ------------------------------
    c, d = api('POST', '/kpi/results/init', TOKU,
               body={'kpiId': kpis['REV_GROWTH']['kpiId'], 'periodId': periods[('ANNUAL', None)]})
    rev_res = d.get('resultId')
    c, d = api('PUT', '/kpi/results/%d' % rev_res, TOKU, body={'figureA': 108, 'figureB': 100, 'notes': 'UAT'})
    ok = d.get('resultPct') == 108 and d.get('score') == 5
    page.evaluate("window._jetApp.navigate('resultEntry', { resultId: %d })" % rev_res)
    page.wait_for_timeout(3500); shot('07')
    record('KPI-07', 'Results', 'Ratio KPI computes and scores live',
           'Enter revenue 108 vs prior 100', 'Result 108% and score 5 preview', ok, d)

    # ---- KPI-08 rubric entry ----------------------------------------------
    c, d = api('POST', '/kpi/results/init', TOKU,
               body={'kpiId': kpis['OPT_PLAN']['kpiId'], 'periodId': periods[('ANNUAL', None)]})
    opt_res = d.get('resultId')
    c, d = api('GET', '/kpi/results/%d' % opt_res, TOKU)
    for cr in d['criteria']:
        c, last = api('PUT', '/kpi/results/%d/criteria/%d' % (opt_res, cr['criterionId']),
                      TOKU, body={'levelNo': 4, 'justification': 'UAT evidence ref'})
    ok = last.get('resultPct') == 80 and last.get('score') == 5
    page.evaluate("window._jetApp.navigate('resultEntry', { resultId: %d })" % opt_res)
    page.wait_for_timeout(3500); shot('08')
    record('KPI-08', 'Results', 'Maturity rubric weighted scoring',
           'Rate all 4 OPT_PLAN criteria at level 4', 'Weighted result 80% and score 5', ok, last)

    # ---- KPI-09 evidence gate ---------------------------------------------
    c, d = api('POST', '/kpi/results/%d/submit' % opt_res, TOKU)
    record('KPI-09', 'Validation', 'Submission blocked without required evidence',
           'Submit OPT_PLAN with no documents', 'HTTP 400 mentioning evidence',
           c == 400 and 'evidence' in str(d), (c, d)); shot('09')

    # ---- KPI-10 oversize upload -------------------------------------------
    c, d = api('PUT', '/kpi/results/%d/docs?file_name=big.bin' % opt_res, TOKU,
               raw=b'x' * (11 * 1024 * 1024))
    record('KPI-10', 'Evidence', 'Upload above MAX_UPLOAD_MB rejected',
           'Upload an 11 MB file (cap 10 MB)', 'HTTP 413', c == 413, c); shot('10')

    # ---- KPI-11 evidence upload + download --------------------------------
    pdf = b'%PDF-1.4 uat evidence\n%%EOF'
    c1, d = api('PUT', '/kpi/results/%d/docs?file_name=uat.pdf&mime_type=application/pdf' % opt_res,
                TOKU, raw=pdf, mime='application/pdf')
    doc_id = d.get('docId')
    c2, d2 = api('GET', '/kpi/results/%d/docs' % opt_res, TOKU)
    c3, _ = api('GET', '/kpi/docs/%d/file' % doc_id, TOKU)
    page.evaluate("window._jetApp.navigate('resultEntry', { resultId: %d })" % opt_res)
    page.wait_for_timeout(3500); shot('11')
    record('KPI-11', 'Evidence', 'Evidence upload, register and download',
           'Upload a PDF then list and download it', 'Doc row created, download 200',
           c1 == 200 and len(d2.get('items', [])) == 1 and c3 == 200, (c1, c2, c3))

    # ---- KPI-12 submit -> workflow ----------------------------------------
    c, d = api('POST', '/kpi/results/%d/submit' % opt_res, TOKU)
    inst = d.get('wfInstanceId')
    record('KPI-12', 'Workflow', 'Submission starts the approval chain',
           'Submit the completed OPT_PLAN result', 'Status SUBMITTED + workflow instance id',
           c == 200 and d.get('status') == 'SUBMITTED' and bool(inst), d); shot('12')

    # ---- KPI-13 result locked ---------------------------------------------
    c, _ = api('PUT', '/kpi/results/%d' % opt_res, TOKU, body={'figureA': 1, 'figureB': 1, 'notes': ''})
    record('KPI-13', 'Workflow', 'Figures locked while under approval',
           'Try to edit the submitted result', 'HTTP 400', c == 400, c); shot('13')

    # ---- KPI-14 worklist task + outcomes ----------------------------------
    c, d = api('GET', '/wf/worklist', TOKA)
    items = [i for i in d.get('items', []) if i.get('module') == 'KPI_MGMT']
    outcomes = sorted(o.get('code') for o in items[0]['outcomes']) if items else []
    page.evaluate("window._jetApp.navigate('myWorklist')"); page.wait_for_timeout(3500); shot('14')
    record('KPI-14', 'Workflow', 'Task in the shared worklist with server outcomes',
           'Open My Worklist as the Section Head reviewer',
           'KPI task listed; outcomes APPROVE/REJECT/RETURN from the outcome set',
           len(items) >= 1 and outcomes == ['APPROVE', 'REJECT', 'RETURN'], outcomes)

    # ---- KPI-15 two-step approval + hooks ---------------------------------
    api('POST', '/wf/tasks/%s/action' % items[0]['id'], TOKA, body={'outcome': 'APPROVE', 'comments': 'UAT step 1'})
    c, d = api('GET', '/wf/worklist', TOKS)
    it2 = [i for i in d.get('items', []) if i.get('module') == 'KPI_MGMT']
    if it2:
        api('POST', '/wf/tasks/%s/action' % it2[0]['id'], TOKS, body={'outcome': 'APPROVE', 'comments': 'UAT step 2'})
    c, d = api('GET', '/kpi/results/%d' % opt_res, TOKU)
    page.evaluate("window._jetApp.navigate('resultEntry', { resultId: %d })" % opt_res)
    page.wait_for_timeout(3500); shot('15')
    record('KPI-15', 'Workflow', 'Two-step chain approves and completes the result',
           'Approve as Section Head then Finance Director', 'Result APPROVED with 3+ history rows',
           d.get('status') == 'APPROVED' and len(d.get('history', [])) >= 3,
           (d.get('status'), len(d.get('history', []))))

    # ---- KPI-16 scorecard reflects approval --------------------------------
    c, d = api('GET', '/kpi/scorecard?year=%d' % YEAR, TOKU)
    sc = {k['code']: k for k in d.get('kpis', [])}
    page.evaluate("window._jetApp.navigate('dashboard')"); page.wait_for_timeout(3000); shot('16')
    record('KPI-16', 'Dashboard', 'Scorecard reflects the approved measurement',
           'Reload the scorecard for %d' % YEAR, 'OPT_PLAN average score 5, overall score present',
           sc.get('OPT_PLAN', {}).get('avgScore') == 5 and d.get('overallScore') is not None,
           (sc.get('OPT_PLAN', {}).get('avgScore'), d.get('overallScore')))

    # ---- KPI-17 briefing book ---------------------------------------------
    import time as _t
    c, d = api('POST', '/kpi/reports/book', TOKA, body={'year': YEAR})
    run = d.get('runId'); status = ''
    for _ in range(30):
        _t.sleep(8)
        c, d = api('GET', '/kpi/reports/%d/status' % run, TOKA)
        status = d.get('status')
        if status in ('SUCCESS', 'FAILED'): break
    c2, _ = api('GET', '/kpi/reports/%d/pdf' % run, TOKA)
    page.evaluate("window._jetApp.navigate('reports')"); page.wait_for_timeout(2500); shot('17')
    record('KPI-17', 'Reports', 'KPI Briefing Book generated and downloadable',
           'Generate the briefing book for %d via the Reports page bridge' % YEAR,
           'Run SUCCESS with a downloadable PDF', status == 'SUCCESS' and c2 == 200, (status, c2))

    # ---- KPI-18 AR / RTL ---------------------------------------------------
    page.locator('.lang-pill button').nth(1).click(); page.wait_for_timeout(3000)
    ok = page.evaluate("document.documentElement.dir") == 'rtl'
    shot('18')
    page.locator('.lang-pill button').nth(0).click(); page.wait_for_timeout(2500)
    record('KPI-18', 'i18n', 'Arabic interface with RTL layout',
           'Switch the language to Arabic (then restore English)',
           'Document flips to RTL with translated labels', ok,
           page.evaluate("document.documentElement.dir"))

    browser.close()

# ---------------------------------------------------------------- artifacts
HEAD = ['Case ID', 'Area', 'Test Case', 'Steps', 'Expected Result', 'Status', 'Actual / Notes']
BRAND = '8C6D1F'

def build_wb(path, with_status):
    wb = Workbook(); ws = wb.active; ws.title = 'UAT'
    ws.append(['Finance KPIs (App 213) — UAT ' + ('Round %d — %s' % (ROUND, DMON) if with_status else 'Test Script')])
    ws['A1'].font = Font(bold=True, size=14, color=BRAND); ws.append([])
    ws.append(HEAD)
    for c in ws[3]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=BRAND)
    for (cid, area, title, steps, exp, st, act) in CASES:
        ws.append([cid, area, title, steps, exp, st if with_status else '', str(act) if with_status else ''])
        row = ws.max_row
        if with_status:
            ws.cell(row=row, column=6).font = Font(bold=True, color='1B7F4B' if st == 'PASS' else 'C7351F')
        for col in range(1, 8):
            ws.cell(row=row, column=col).alignment = Alignment(vertical='top', wrap_text=True)
    for col, w in zip('ABCDEFG', [10, 12, 34, 40, 40, 9, 30]):
        ws.column_dimensions[col].width = w
    wb.save(path)

build_wb(os.path.join(ROOT, 'UAT_KPI_TestScript.xlsx'), False)
build_wb(os.path.join(RDIR, 'UAT_KPI_%s-01.xlsx' % DMON), True)

doc = Document()
h = doc.add_heading('Finance KPIs (App 213) — UAT Results', 0)
doc.add_paragraph('Round %d · %s · Automated runner: assessment-3/phase4/tests/uat_run_kpi.py' % (ROUND, DMON))
npass = sum(1 for c in CASES if c[5] == 'PASS')
s = doc.add_paragraph(); r = s.add_run('Summary: %d / %d cases PASS' % (npass, len(CASES)))
r.bold = True; r.font.color.rgb = RGBColor(0x1B, 0x7F, 0x4B) if npass == len(CASES) else RGBColor(0xC7, 0x35, 0x1F)
tbl = doc.add_table(rows=1, cols=5); tbl.style = 'Light Grid Accent 1'
for i, htxt in enumerate(['Case', 'Area', 'Test Case', 'Status', 'Actual / Notes']):
    tbl.rows[0].cells[i].paragraphs[0].add_run(htxt).bold = True
for (cid, area, title, steps, exp, st, act) in CASES:
    row = tbl.add_row().cells
    row[0].text = cid; row[1].text = area; row[2].text = title; row[3].text = st; row[4].text = str(act)[:120]
doc.add_heading('Evidence', 1)
for (cid, area, title, steps, exp, st, act) in CASES:
    png = os.path.join(EDIR, 'case_%s.png' % cid.split('-')[1])
    if os.path.exists(png):
        doc.add_paragraph('%s — %s (%s)' % (cid, title, st))
        doc.add_picture(png, width=Inches(6.2))
doc.save(os.path.join(RDIR, 'UAT_KPI_Results_%s-01.docx' % DMON))

print('-' * 50)
print('UAT %d/%d PASS — package in %s' % (npass, len(CASES), RDIR))
sys.exit(0 if npass == len(CASES) else 1)
