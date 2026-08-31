#!/usr/bin/env python3
"""Load a Fusion/EBS segment mapping Excel into PROD.DCT_GL_EBS_MAP.

Rerunnable: MERGE keyed on (segment_type, ebs_value) -- existing rows are
updated, new rows inserted, nothing deleted (deactivate via the GL app).

Expected sheet layout (header row 1, like "COA Account_Mapping.xlsx"):
    EBS_VALUE | FUSION_VALUE | PARENT_CHILD | FUSION_VALUE_DESCRIPTION

Usage (dev VM, wallet at /opt/oracle-wallet/Wallet_prod):
    python3 load_coa_map.py --segment ACCOUNT --file "COA Account_Mapping.xlsx"
    python3 load_coa_map.py --segment APPROPRIATION --file "Future2_Mapping.xlsx"

Note 2026-07-30: the APPROPRIATION map turned out to be IDENTITY (user rule:
Future2 IS the appropriation code, just unpadded in EBS) — it is seeded by
db/v2/110 section 110.5b and no mapping file is needed. This loader remains
for corrections/overrides only (ebs_value in the 6-digit padded form).
"""
import argparse
import getpass
import os
import sys

import oracledb
from openpyxl import load_workbook

WALLET = os.environ.get("PROD_WALLET", "/opt/oracle-wallet/Wallet_prod")

MERGE = """
MERGE INTO prod.dct_gl_ebs_map t
USING (SELECT :seg AS segment_type, :ebs AS ebs_value FROM dual) s
   ON (t.segment_type = s.segment_type AND t.ebs_value = s.ebs_value)
WHEN NOT MATCHED THEN INSERT
     (segment_type, ebs_value, fusion_value, parent_child, description, created_by)
VALUES (:seg, :ebs, :fus, :pc, :descr, :who)
WHEN MATCHED THEN UPDATE SET
     t.fusion_value = :fus,
     t.parent_child = :pc,
     t.description  = :descr,
     t.updated_by   = :who,
     t.updated_at   = SYSTIMESTAMP
"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--segment", required=True, choices=["ACCOUNT", "APPROPRIATION"])
    ap.add_argument("--file", required=True)
    ap.add_argument("--sheet", default=None, help="sheet name (default: first)")
    ap.add_argument("--password", default=os.environ.get("PROD_ADMIN_PWD"))
    args = ap.parse_args()

    pwd = args.password or getpass.getpass("ADMIN password: ")
    wb = load_workbook(args.file, read_only=True)
    ws = wb[args.sheet] if args.sheet else wb[wb.sheetnames[0]]

    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or r is None or r[0] is None or r[1] is None:
            continue
        pc = (str(r[2]).strip().upper() if len(r) > 2 and r[2] else "CHILD")
        desc = (str(r[3]).strip()[:400] if len(r) > 3 and r[3] else None)
        rows.append({"seg": args.segment, "ebs": str(r[0]).strip(),
                     "fus": str(r[1]).strip(), "pc": pc, "descr": desc,
                     "who": "SETUP"})
    if not rows:
        sys.exit("no data rows found")
    print(f"{len(rows)} rows parsed from {args.file}")

    conn = oracledb.connect(user="ADMIN", password=pwd, dsn="prod_low",
                            config_dir=WALLET, wallet_location=WALLET,
                            wallet_password=pwd)
    cur = conn.cursor()
    cur.executemany(MERGE, rows)
    conn.commit()
    cur.execute("""SELECT segment_type, COUNT(*), SUM(CASE WHEN is_active='Y' THEN 1 ELSE 0 END)
                   FROM prod.dct_gl_ebs_map GROUP BY segment_type""")
    for seg, n, act in cur.fetchall():
        print(f"  {seg}: {n} rows ({act} active)")
    conn.close()


if __name__ == "__main__":
    main()
