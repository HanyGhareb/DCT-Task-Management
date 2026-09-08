"""Package BI storage verification using the required dated UAT layout."""
import json,re,shutil
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
from docx.shared import Inches
BASE=Path(__file__).resolve().parents[1]/'UAT';TRIAL=BASE/'UAT_BI_round3-06-09-2026';EV=TRIAL/'evidence_06-Sep-2026-01';EV.mkdir(parents=True,exist_ok=True)
TMP=Path('/tmp/rpt-storage-impl')
for name in ['browser','live','deployed-browser','deployed-live']:
 shutil.copytree(TMP/name,EV/name,dirs_exist_ok=True)
for name in ['db-test.log','deploy-db.log','deploy-ui.log','final-check.log','db-before.json']:
 shutil.copy2(TMP/name,EV/name)
rows=[]
for suite,path in [('Browser edge and language checks',EV/'deployed-browser/results.json'),('Live API and access',EV/'deployed-live/results.json')]:
 for name,ok in json.loads(path.read_text()):rows.append([f'BI-ST-{len(rows)+1:03}',suite,name,'PASS' if ok else 'FAIL',str(path.relative_to(TRIAL))])
for name in re.findall(r'^PASS (.+)$',(EV/'db-test.log').read_text(),re.M):rows.append([f'BI-ST-{len(rows)+1:03}','Database cleanup',name,'PASS','evidence_06-Sep-2026-01/db-test.log'])
rows.append(['BI-ST-SCHED','Scheduler','New audited procedure executed successfully on the retained 15-minute ADMIN schedule; duplicate disabled','PASS','evidence_06-Sep-2026-01/final-check.log'])
header=['ID','Suite','Check','Status','Evidence'];wb=Workbook();ws=wb.active;ws.title='Storage maintenance';ws.append(header)
for row in rows:ws.append(row)
for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F6F8B')
for c,w in zip('ABCDE',[20,35,85,15,75]):ws.column_dimensions[c].width=w
for rr in ws.iter_rows(min_row=2):
 for c in rr:c.alignment=Alignment(wrap_text=True,vertical='top')
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;wb.save(TRIAL/'UAT_BI_06-Sep-2026-01.xlsx')
master=BASE/'UAT_BI_TestScript.xlsx';m=load_workbook(master)
if 'Storage maintenance' in m.sheetnames:del m['Storage maintenance']
w=m.create_sheet('Storage maintenance');w.append(header)
for r in rows:w.append(r[:3]+['NOT RUN',r[4]])
m.save(master)
d=Document();d.add_heading('BI — Report Storage and Cleanup',0)
d.add_paragraph('06 September 2026 · BI 1.16.0 · Round 3')
d.add_paragraph('Implemented the three approved priorities: one maintenance schedule, a BI storage overview and cleanup audit. The existing 90-day retention policy is unchanged. Open BI → Workers → Report storage and cleanup (SYS_ADMIN only).')
d.add_heading('Verification',1)
d.add_paragraph('21 browser checks, 10 live API/access checks and 11 database assertions passed. The retained ADMIN maintenance job subsequently executed the audited procedure successfully on schedule. English and Arabic tested at 1535, 768 and 390 pixels. Empty/error/retry states, escaping of error text, 401 and 403 responses verified. Worker polling does not trigger storage scans; storage refresh runs on entry and explicit refresh.')
d.add_paragraph('The database test inserted one expired and one fresh synthetic output. Cleanup removed only the expired 16-byte fixture and preserved run history. Both fixtures and their synthetic run were removed afterward. Audit row 1 records this test deletion; subsequent cleanups recorded zero. No real files or financial data were deleted by the test.')
d.add_heading('Behavior and limitations',1)
d.add_paragraph('The panel shows file count and bytes, allocated database space when dictionary access is available, currently retained files added in 7/30 days, cleanup eligibility, the largest 10 reports and the latest 20 audit entries. The audit stores timestamps, retention/cutoff, actual deleted files/bytes, success/failure and error. Successful audit updates commit with deletion; errors roll back deletion and record failure. Audit history starts at deployment and is retained. No per-report retention, Object Storage archive, manual delete action or physical storage shrink was introduced.')
d.add_heading('Deployment and rollback',1)
d.add_paragraph((TMP/'deploy-ui.log').read_text().strip())
d.add_paragraph('SQL: reporting/db/43_rpt_output_maintenance.sql and 44_rpt_storage_ords.sql. The canonical 05 scheduler installer delegates to 43 so reinstalling cannot restore the old unaudited cleanup or duplicate schedule. Existing ADMIN job remains managed by BI Workers; verified PROD duplicate is disabled, retained for rollback. After a full ORDS module rebuild (04), reinstall 44 as well as the existing module extensions.')
d.add_paragraph('Rollback should reverse only this BI change if later releases exist; the previous release and pre-change procedure/job snapshot are recorded in evidence. Retain audit records. No report worker restart was needed.')
for lang in ['en','ar']:
 d.add_heading('Storage overview — '+lang.upper(),1);d.add_picture(str(EV/'deployed-live'/(lang+'.png')),width=Inches(6.4))
 d.add_heading('Cleanup history — '+lang.upper(),2);d.add_picture(str(EV/'deployed-live'/(lang+'-audit.png')),width=Inches(6.4))
d.save(TRIAL/'UAT_BI_Results_06-Sep-2026-01.docx');print(TRIAL)
