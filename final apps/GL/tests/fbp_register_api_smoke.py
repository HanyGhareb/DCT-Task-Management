#!/usr/bin/env python3
"""GL /gl/butil/fbpxlsx API smoke — the FBP - Projects Budget Utilization bridge.

FBP_BUTIL_REGISTER (reporting/db/45) is the Budget Utilization Register
re-issued for the FBP distribution: sheet 1 carries a FIXED 21-column layout
(Task Name in place of Task Number; the red-marked columns of the sample
docs/Reports/FMR/FBP-Budget_Utilization_Register_2026.xlsx removed), all other
worksheets are a verbatim copy. This smoke enqueues a run through the GL
bridge (GL/db/52), waits for a worker to render it, downloads the workbook and
verifies the fixed sheet-1 layout plus the guardrails (auth, validation, run
scoping — and that the page's Manage-Columns sheetcols is IGNORED).

Run: python fbp_register_api_smoke.py            (defaults to the webtier)
     GL_BASE=http://localhost:8210 python fbp_register_api_smoke.py
     GL_PERIOD=09-2026 python fbp_register_api_smoke.py
"""
import io
import json
import os
import re
import ssl
import time
import urllib.error
import urllib.request

BASE = os.environ.get('GL_BASE', 'https://129.151.159.189')
AUTH = '/root/DCT-Task-Management/final apps/Admin/Jet/js/services/authService.js'
YEAR = int(os.environ.get('GL_YEAR', '2026'))
PERIOD = os.environ.get('GL_PERIOD', time.strftime('%m-') + str(YEAR))
OUT = os.environ.get('GL_XLSX_OUT', '/tmp/fbp_projects_budget_utilization_smoke.xlsx')
CTX = ssl.create_default_context()
CTX.check_hostname = False
CTX.verify_mode = ssl.CERT_NONE

# the FBP agreed sheet-1 layout — 21 columns, in order (Task Name, not Number)
FBP_COLS = ['Budget Combination', 'Sector', 'Department', 'Cost Centre',
            'Project Number', 'Project Name', 'Task Name', 'Account Number',
            'Appropriation Code', 'Chapter', 'Dct Program Code',
            'Expenditure Type', 'Annual Budget', 'Actual Ap', 'Actual Grn',
            'Actual Total', 'Commitment Pr', 'Obligation Po',
            'Open Encumbrance', 'Fund Available', 'Comments']
REMOVED = ['Task Number', 'Ebs Account', 'Appropriation Name',
           'Dct Program Name', 'Ytd Budget', 'Budget Utilization Pct',
           'Budget Variance', 'Budget Status', 'Plan Annual', 'Plan Ytd',
           'Revised Plan Annual', 'Plan Utilization Pct', 'Plan Variance',
           'Plan Status', 'Plan Coverage Pct', 'Plan Coverage Status',
           'Utilization Pct', 'Fund Movement Count', 'Fund Movement Amount']

ok = fail = 0


def ck(name, cond, extra=''):
    global ok, fail
    print(('  PASS ' if cond else '  FAIL ') + name + ('  ' + str(extra) if extra else ''))
    if cond:
        ok += 1
    else:
        fail += 1


def login():
    u, p = re.findall(r"username:\s*'([^']+)',\s*password:\s*'([^']+)'", open(AUTH).read())[0]
    r = urllib.request.Request(BASE + '/ords/admin/dct/auth/login',
                               data=json.dumps({'username': u, 'password': p}).encode(),
                               method='POST', headers={'Content-Type': 'application/json'})
    return json.loads(urllib.request.urlopen(r, timeout=60, context=CTX).read())['sessionId']


def api(tok, path, method='GET', body=None):
    hdr = {'Authorization': 'Bearer ' + tok}
    data = None
    if body is not None:
        hdr['Content-Type'] = 'application/json'
        data = json.dumps(body).encode()
    r = urllib.request.Request(BASE + '/ords/admin/gl' + path, method=method,
                               data=data, headers=hdr)
    return json.loads(urllib.request.urlopen(r, timeout=180, context=CTX).read())


def status_of(tok, path, method='GET', body=None, with_tok=True):
    hdr = {'Authorization': 'Bearer ' + tok} if with_tok else {}
    data = None
    if body is not None:
        hdr['Content-Type'] = 'application/json'
        data = json.dumps(body).encode()
    r = urllib.request.Request(BASE + '/ords/admin/gl' + path, method=method,
                               data=data, headers=hdr)
    try:
        urllib.request.urlopen(r, timeout=60, context=CTX)
        return 200
    except urllib.error.HTTPError as e:
        return e.code


def main():
    tok = login()
    ck('login', bool(tok))

    # guardrails
    ck('POST without session = 401',
       status_of('', '/butil/fbpxlsx', 'POST', {'year': YEAR}, with_tok=False) == 401)
    ck('POST without year = 400',
       status_of(tok, '/butil/fbpxlsx', 'POST', {}) == 400)
    ck('POST bad period = 400',
       status_of(tok, '/butil/fbpxlsx', 'POST',
                 {'year': YEAR, 'period': '13-' + str(YEAR)}) == 400)
    ck('GET unknown run = 404', status_of(tok, '/butil/fbpxlsx/999999999') == 404)

    # enqueue with the default page scope; sheetcols MUST be ignored (fixed layout)
    d = api(tok, '/butil/fbpxlsx', 'POST',
            {'year': YEAR, 'period': PERIOD,
             'projecttype': 'DCT OPEX Project Type',
             'sheetcols': 'project_number,task_number'})
    run_id = d.get('runId')
    ck('enqueue -> runId', bool(run_id), 'run #%s (period %s)' % (run_id, PERIOD))

    # a BUDGET_UTIL_REGISTER status route must NOT see this run (report-code scoping)
    ck('run invisible to /butil/xlsx/:id', status_of(tok, '/butil/xlsx/%s' % run_id) == 404)

    st = {}
    for _ in range(75):                     # ~7.5 min ceiling
        time.sleep(6)
        try:
            st = api(tok, '/butil/fbpxlsx/%s' % run_id)
        except Exception:
            continue
        if st.get('status') in ('SUCCESS', 'FAILED'):
            break
    ck('run finished SUCCESS', st.get('status') == 'SUCCESS', st.get('error') or st.get('status'))
    ck('run hasFile', bool(st.get('hasFile')))

    r = urllib.request.Request(BASE + '/ords/admin/gl/butil/fbpxlsx/%s/file' % run_id,
                               headers={'Authorization': 'Bearer ' + tok})
    xls = urllib.request.urlopen(r, timeout=300, context=CTX).read()
    ck('XLSX downloads', xls[:2] == b'PK', '%d bytes' % len(xls))
    with open(OUT, 'wb') as f:
        f.write(xls)
    print('  workbook saved to ' + OUT)

    # ── workbook layout verification ──
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xls), read_only=True)
    names = wb.sheetnames
    ck('8 worksheets', len(names) == 8, names)
    ck('sheet 1 = Budget Utilization Lines', 'Budget Utilization Lines' in names[0], names[0])
    ck('sheet 8 = Fund Movement', 'Fund Movement' in names[-1], names[-1])

    def headers_of(ws, probe):
        for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
            vals = [str(v).strip() for v in row if v not in (None, '')]
            if probe in vals:
                return vals
        return []

    h1 = headers_of(wb[names[0]], 'Budget Combination')
    ck('sheet 1 = the fixed 21-column FBP layout', h1 == FBP_COLS,
       h1 if h1 != FBP_COLS else '21 columns, in order')
    ck('sheet 1 has Task Name, no Task Number',
       'Task Name' in h1 and 'Task Number' not in h1)
    gone = [c for c in REMOVED if c in h1]
    ck('all red-marked columns removed from sheet 1', not gone, gone or '19 removed')

    # untouched copies: AP sheet keeps Task Number; pending sheet unchanged
    h2 = headers_of(wb[names[1]], 'Sector')
    ck('sheet 2 (AP) untouched — keeps Task Number', 'Task Number' in h2)
    h6 = headers_of(wb[names[5]], 'Doc Type')
    ck('sheet 6 (Pending) untouched', 'Pending With' in h6 and 'Gl Combination' in h6)
    wb.close()

    print('\n%d passed, %d failed' % (ok, fail))
    raise SystemExit(1 if fail else 0)


if __name__ == '__main__':
    main()
