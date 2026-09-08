"""Create dated live verification workbook/results from Budget Status smoke outputs."""
import json
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
from docx.shared import Inches

base=Path(__file__).resolve().parents[1]/'UAT'
trial=base/'UAT_GL_round2-06-09-2026'
ev=trial/'evidence_06-Sep-2026-01'
sources=[('Dashboard regression','FD-BROWSER',trial/'browser-results.json'),
         ('Report API and totals','FD-REPORT',ev/'reports/results.json'),
         ('Print menu / 1.113.0','FD-PRINT',ev/'print-ui/results.json'),
         ('Latest release / 1.115.0','FD-LATEST',ev/'latest-release/results.json')]
rows=[];counts=[]
for label,prefix,path in sources:
    checks=json.loads(path.read_text());counts.append((label,sum(ok for _,ok in checks),len(checks)))
    for i,(name,ok) in enumerate(checks,1):
        rows.append([f'{prefix}-{i:03}',name,label,'PASS' if ok else 'FAIL',str(path.relative_to(trial))])
rows.extend([
 ['FD-DB-01','SQLcl scripts compile/execute; report definition JSON valid','Production database','PASS','docs/deployment-notes.md'],
 ['FD-DB-02','Three report handlers installed; no report recipients/deliveries','Production database','PASS','docs/deployment-notes.md'],
 ['FD-ACCESS-01','Live 403 for a separate user without report privilege','Separate user not provisioned','NOT RUN','403 recovery tested locally; deployed privilege guards reviewed'],
 ['FD-ACCESS-02','Live second-user report ownership isolation','Separate user not provisioned','NOT RUN','Requesting-user predicates reviewed in deployed status/file handlers']])
wb=Workbook();ws=wb.active;ws.title='Live verification';ws.append(['ID','Check','Suite / environment','Status','Evidence / notes'])
for row in rows:ws.append(row)
for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='3C6D5D')
for col,width in zip('ABCDE',[24,85,32,16,75]):ws.column_dimensions[col].width=width
for row in ws.iter_rows(min_row=2):
    for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
wb.save(trial/'UAT_GL_06-Sep-2026-01.xlsx')
master=base/'UAT_GL_TestScript.xlsx'
m=load_workbook(master)
if 'Live Budget Status' not in m.sheetnames:
    target=m.create_sheet('Live Budget Status');target.append(['ID','Check','Suite / environment','Status','Evidence / notes'])
    for r in rows:target.append(r[:3]+['NOT RUN']+[r[4]])
    for col,width in zip('ABCDE',[24,85,32,16,75]):target.column_dimensions[col].width=width
    target.freeze_panes='A2';m.save(master)
d=Document();d.add_heading('Budget Status — Production Verification',0)
d.add_paragraph('06 September 2026 · Round 2 · Deployment approved by the user. Initial feature version 1.113.0; retained and rechecked in GL 1.115.0.')
d.add_heading('Outcome',1)
d.add_paragraph('Compact Search and Presentation layout is live. Show Summary starts unchecked. Print produces PDF and PowerPoint using the reporting service, with current filters and presentation and every matching department. PowerPoint contains page artwork images. No previous-period comparison was added.')
t=d.add_table(rows=1,cols=2);t.style='Light Shading Accent 1';t.rows[0].cells[0].text='Suite';t.rows[0].cells[1].text='Passed'
for label,passed,total in counts:
    c=t.add_row().cells;c[0].text=label;c[1].text=f'{passed}/{total}'
d.add_paragraph('The later-release check repeats the Print/layout suite to verify concurrent GL changes preserved this feature. Database scripts installed successfully; all three reporting workers are active. Six report API jobs (1002–1007) and actual browser downloads succeeded. No recipients or delivery records were created. Live 401/400/404 passed; 403 was tested using local fixtures. Separate-user privilege/ownership tests were not run; deployed guards were reviewed.')
d.add_heading('Latest live search',1);d.add_picture(str(ev/'latest-release/search-live.png'),width=Inches(6.4))
d.add_heading('Latest live Arabic search',1);d.add_picture(str(ev/'latest-release/search-live-ar.png'),width=Inches(6.4))
d.add_heading('Deployment and rollback',1)
d.add_paragraph('GL-only release 20260905201058, previously 20260904081416. Later GL 1.115.0 is at release 20260905202533 and includes these changes. Do not roll back later unrelated changes: reverse only this feature’s hunks. Worker backups are /opt/rpt-worker/backups/budget-status-20260906 on vm180/181/182. Full details are in GL/docs/budget-status-review.md and deployment-notes.md.')
d.add_heading('Evidence',1)
d.add_paragraph('The dated workbook contains every check and links to machine-readable results. evidence_06-Sep-2026-01 contains browser screenshots, real PDF/PPT downloads, report run metadata and the dashboard cube used for reconciliation. These round-2 files contain production financial data; round-1 samples contain synthetic data.')
d.save(trial/'UAT_GL_Results_06-Sep-2026-01.docx')
print('Created live workbook and Word results:',trial)
