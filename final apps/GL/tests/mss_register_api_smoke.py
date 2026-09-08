#!/usr/bin/env python3
"""GL /gl/butil/mssxlsx API smoke — the MSS - Projects Budget Utilization bridge.

MSS_BUTIL_REGISTER (reporting/db/46) is the Budget Utilization Register
re-issued for the MSS distribution per the marked-up sample
docs/Reports/FMR/MSS Budget_Utilization_Register_2026.xlsx: sheet 1 = FIXED
26-column layout (Task Number kept + Task Name added; the 14 red-marked
columns removed), Requester added on sheets 2-6 (AP = the matched PO line's
requester, GRN/PO = the PO distribution requestor, PR = the requisition
requester, Pending = per document) and a Task Name column beside EVERY
Task Number (all sheets, 2026-09-08 user request). This smoke
enqueues a run through the GL bridge (GL/db/53), waits for a worker, downloads
the workbook and verifies every sheet's layout + requester fill + guardrails.

Run: python mss_register_api_smoke.py            (defaults to the webtier)
     GL_PERIOD=09-2026 python mss_register_api_smoke.py
"""
import io
import json
import os
import re
import ssl
import time
import urllib.error
import urllib.request

BASE = os.environ.get('GL_BASE', 'https://129.151.159.189')
AUTH = '/root/DCT-Task-Management/final apps/Admin/Jet/js/services/authService.js'
YEAR = int(os.environ.get('GL_YEAR', '2026'))
PERIOD = os.environ.get('GL_PERIOD', time.strftime('%m-') + str(YEAR))
OUT = os.environ.get('GL_XLSX_OUT', '/tmp/mss_projects_budget_utilization_smoke.xlsx')
CTX = ssl.create_default_context()
CTX.check_hostname = False
CTX.verify_mode = ssl.CERT_NONE

# the MSS agreed sheet-1 layout — 26 columns, in order
MSS_S1 = ['Budget Combination', 'Sector', 'Department', 'Cost Centre',
          'Project Number', 'Project Name', 'Task Number', 'Task Name',
          'Account Number', 'Ebs Account', 'Chapter', 'Expenditure Type',
          'Annual Budget', 'Budget Utilization Pct', 'Budget Variance',
          'Actual Ap', 'Actual Grn', 'Actual Total', 'Commitment Pr',
          'Obligation Po', 'Open Encumbrance', 'Fund Available',
          'Utilization Pct', 'Fund Movement Count', 'Fund Movement Amount',
          'Comments']
S1_REMOVED = ['Appropriation Code', 'Appropriation Name', 'Dct Program Code',
              'Dct Program Name', 'Ytd Budget', 'Budget Status', 'Plan Annual',
              'Plan Ytd', 'Revised Plan Annual', 'Plan Utilization Pct',
              'Plan Variance', 'Plan Status', 'Plan Coverage Pct',
              'Plan Coverage Status']

ok = fail = 0


def ck(name, cond, extra=''):
    global ok, fail
    print(('  PASS ' if cond else '  FAIL ') + name + ('  ' + str(extra) if extra else ''))
    if cond:
        ok += 1
    else:
        fail += 1


def login():
    u, p = re.findall(r"username:\s*'([^']+)',\s*password:\s*'([^']+)'", open(AUTH).read())[0]
    r = urllib.request.Request(BASE + '/ords/admin/dct/auth/login',
                               data=json.dumps({'username': u, 'password': p}).encode(),
                               method='POST', headers={'Content-Type': 'application/json'})
    return json.loads(urllib.request.urlopen(r, timeout=60, context=CTX).read())['sessionId']


def api(tok, path, method='GET', body=None):
    hdr = {'Authorization': 'Bearer ' + tok}
    data = None
    if body is not None:
        hdr['Content-Type'] = 'application/json'
        data = json.dumps(body).encode()
    r = urllib.request.Request(BASE + '/ords/admin/gl' + path, method=method,
                               data=data, headers=hdr)
    return json.loads(urllib.request.urlopen(r, timeout=180, context=CTX).read())


def status_of(tok, path, method='GET', body=None, with_tok=True):
    hdr = {'Authorization': 'Bearer ' + tok} if with_tok else {}
    data = None
    if body is not None:
        hdr['Content-Type'] = 'application/json'
        data = json.dumps(body).encode()
    r = urllib.request.Request(BASE + '/ords/admin/gl' + path, method=method,
                               data=data, headers=hdr)
    try:
        urllib.request.urlopen(r, timeout=60, context=CTX)
        return 200
    except urllib.error.HTTPError as e:
        return e.code


def main():
    tok = login()
    ck('login', bool(tok))

    # guardrails
    ck('POST without session = 401',
       status_of('', '/butil/mssxlsx', 'POST', {'year': YEAR}, with_tok=False) == 401)
    ck('POST without year = 400',
       status_of(tok, '/butil/mssxlsx', 'POST', {}) == 400)
    ck('POST bad period = 400',
       status_of(tok, '/butil/mssxlsx', 'POST',
                 {'year': YEAR, 'period': '13-' + str(YEAR)}) == 400)
    ck('GET unknown run = 404', status_of(tok, '/butil/mssxlsx/999999999') == 404)

    # enqueue; sheetcols MUST be ignored (fixed layout)
    d = api(tok, '/butil/mssxlsx', 'POST',
            {'year': YEAR, 'period': PERIOD,
             'projecttype': 'DCT OPEX Project Type|MSS OPEX Project Type',
             'sheetcols': 'project_number,task_number'})
    run_id = d.get('runId')
    ck('enqueue -> runId', bool(run_id), 'run #%s (period %s)' % (run_id, PERIOD))
    ck('run invisible to /butil/xlsx/:id', status_of(tok, '/butil/xlsx/%s' % run_id) == 404)
    ck('run invisible to /butil/fbpxlsx/:id', status_of(tok, '/butil/fbpxlsx/%s' % run_id) == 404)

    st = {}
    for _ in range(75):                     # ~7.5 min ceiling
        time.sleep(6)
        try:
            st = api(tok, '/butil/mssxlsx/%s' % run_id)
        except Exception:
            continue
        if st.get('status') in ('SUCCESS', 'FAILED'):
            break
    ck('run finished SUCCESS', st.get('status') == 'SUCCESS', st.get('error') or st.get('status'))
    ck('run hasFile', bool(st.get('hasFile')))

    r = urllib.request.Request(BASE + '/ords/admin/gl/butil/mssxlsx/%s/file' % run_id,
                               headers={'Authorization': 'Bearer ' + tok})
    xls = urllib.request.urlopen(r, timeout=300, context=CTX).read()
    ck('XLSX downloads', xls[:2] == b'PK', '%d bytes' % len(xls))
    with open(OUT, 'wb') as f:
        f.write(xls)
    print('  workbook saved to ' + OUT)

    # ── workbook layout verification ──
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xls), read_only=True)
    names = wb.sheetnames
    ck('8 worksheets', len(names) == 8, names)

    def sheet_rows(ws, probe):
        """header list + data rows (list of dicts) located by the probe header"""
        hdr, data = [], []
        for row in ws.iter_rows(values_only=True):
            vals = [v for v in row]
            svals = [str(v).strip() for v in vals if v not in (None, '')]
            if not hdr and probe in svals:
                hdr = [str(v).strip() if v not in (None, '') else '' for v in vals]
                continue
            if hdr and any(v not in (None, '') for v in vals):
                data.append(dict(zip(hdr, vals)))
        return [h for h in hdr if h], data

    def fill_rate(rows, col):
        n = sum(1 for r in rows if r.get(col) not in (None, ''))
        return n, len(rows)

    # sheet 1: fixed 26 columns
    h1, _ = sheet_rows(wb[names[0]], 'Budget Combination')
    ck('sheet 1 = the fixed 26-column MSS layout', h1 == MSS_S1,
       h1 if h1 != MSS_S1 else '26 columns, in order')
    ck('sheet 1 keeps Task Number AND adds Task Name',
       'Task Number' in h1 and 'Task Name' in h1)
    gone = [c for c in S1_REMOVED if c in h1]
    ck('all 14 red-marked columns removed from sheet 1', not gone, gone or '14 removed')

    # sheet 2 (AP): Requester after Organization; filled only on PO-matched rows
    h2, d2 = sheet_rows(wb[names[1]], 'Sector')
    ck('sheet 2 has Requester after Organization',
       'Requester' in h2 and h2.index('Requester') == h2.index('Organization') + 1, h2[:6])
    ck('sheet 2 has Task Name after Task Number',
       'Task Name' in h2 and h2.index('Task Name') == h2.index('Task Number') + 1)
    n2t, t2t = fill_rate(d2, 'Task Name')
    ck('sheet 2 Task Name filled', t2t > 0 and n2t / t2t > 0.5, '%d/%d' % (n2t, t2t))
    direct = [r for r in d2 if r.get('Line Type') == 'Direct']
    matched = [r for r in d2 if r.get('Line Type') not in (None, '', 'Direct', 'Cost Adjustment')]
    n_dir, t_dir = fill_rate(direct, 'Requester')
    n_var, t_var = fill_rate(matched, 'Requester')
    # a Direct row CAN legitimately resolve a requester when its own invoice
    # also carries PO-matched dists netting below the display threshold — only
    # a tiny fraction; anything more = the invoice-number map is fanning out
    ck('sheet 2 Direct rows essentially requester-free (<1%)',
       t_dir > 0 and n_dir / t_dir < 0.01, '%d/%d' % (n_dir, t_dir))
    ck('sheet 2 PO-matched variance rows carry the PO requester',
       t_var == 0 or n_var > 0, '%d/%d variance rows filled' % (n_var, t_var))

    # sheet 3 (GRN): Requester = the receipt PO's requestor
    h3, d3 = sheet_rows(wb[names[2]], 'Po Number')
    ck('sheet 3 has Requester after Organization',
       'Requester' in h3 and h3.index('Requester') == h3.index('Organization') + 1)
    ck('sheet 3 has Task Name after Task Number',
       'Task Name' in h3 and h3.index('Task Name') == h3.index('Task Number') + 1)
    n3t, t3t = fill_rate(d3, 'Task Name')
    ck('sheet 3 Task Name filled', t3t > 0 and n3t / t3t > 0.5, '%d/%d' % (n3t, t3t))
    n3, t3 = fill_rate(d3, 'Requester')
    ck('sheet 3 Requester well-filled', t3 > 0 and n3 / t3 > 0.9, '%d/%d' % (n3, t3))

    # sheet 4 (Open PO): Requester + Task Name
    h4, d4 = sheet_rows(wb[names[3]], 'Funds Status')
    ck('sheet 4 has Requester after Organization',
       'Requester' in h4 and h4.index('Requester') == h4.index('Organization') + 1)
    ck('sheet 4 has Task Name after Task Number',
       'Task Name' in h4 and h4.index('Task Name') == h4.index('Task Number') + 1)
    n4, t4 = fill_rate(d4, 'Requester')
    ck('sheet 4 Requester well-filled', t4 > 0 and n4 / t4 > 0.9, '%d/%d' % (n4, t4))

    # sheet 5 (Open PR): Requester + Task Name
    h5, d5 = sheet_rows(wb[names[4]], 'Pr Number')
    ck('sheet 5 has Requester after Organization',
       'Requester' in h5 and h5.index('Requester') == h5.index('Organization') + 1)
    ck('sheet 5 has Task Name after Task Number',
       'Task Name' in h5 and h5.index('Task Name') == h5.index('Task Number') + 1)
    n5, t5 = fill_rate(d5, 'Requester')
    ck('sheet 5 Requester well-filled', t5 > 0 and n5 / t5 > 0.9, '%d/%d' % (n5, t5))

    # sheet 6 (Pending): Requester after Preparer Buyer
    h6, d6 = sheet_rows(wb[names[5]], 'Doc Type')
    ck('sheet 6 has Requester after Preparer Buyer',
       'Requester' in h6 and h6.index('Requester') == h6.index('Preparer Buyer') + 1)
    ck('sheet 6 has Task Name after Task Number',
       'Task Name' in h6 and h6.index('Task Name') == h6.index('Task Number') + 1)
    n6, t6 = fill_rate(d6, 'Requester')
    ck('sheet 6 Requester filled', t6 == 0 or n6 > 0, '%d/%d' % (n6, t6))

    # sheets 7-8: Task Name beside Task Number (no Requester on sheet 8)
    ck('sheet 7 = Comments - Other Levels', 'Comments' in names[6], names[6])
    h7, _ = sheet_rows(wb[names[6]], 'Level Code')
    ck('sheet 7 has Task Name after Task Number',
       'Task Name' in h7 and h7.index('Task Name') == h7.index('Task Number') + 1, h7)
    h8, d8 = sheet_rows(wb[names[7]], 'Transaction Num')
    ck('sheet 8 has Task Name after Task Number',
       'Task Name' in h8 and h8.index('Task Name') == h8.index('Task Number') + 1)
    n8t, t8t = fill_rate(d8, 'Task Name')
    ck('sheet 8 Task Name filled (native PBT column)',
       t8t == 0 or n8t / t8t > 0.5, '%d/%d' % (n8t, t8t))
    ck('sheet 8 (Fund Movement) has no Requester', 'Requester' not in h8)
    wb.close()

    print('\n%d passed, %d failed' % (ok, fail))
    raise SystemExit(1 if fail else 0)


if __name__ == '__main__':
    main()
