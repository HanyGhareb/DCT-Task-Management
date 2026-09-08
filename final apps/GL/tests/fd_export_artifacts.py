"""Open real PDF/PPT artifacts after rendering all layouts in both languages."""
import copy,json
from io import BytesIO
from pathlib import Path
import fitz
from pptx import Presentation
from fd_report_test import fixture,render_pdf,render_fd
OUT=Path('/tmp/fd-deep/artifacts');OUT.mkdir(parents=True,exist_ok=True);results=[]
def check(name,ok):
 results.append([name,bool(ok)]);print(('PASS ' if ok else 'FAIL ')+name,flush=True)
for lang in ['en','ar']:
 for layout in ['tiles','rows','map']:
  ctx=fixture();ctx['params'].update(lang=lang,presentation=layout)
  for r in ctx['sections'][0]['data']:
   r['department_name']=r['department']+' '+('Financial Planning and Strategy Coordination ' * 5)
  key=lang+'-'+layout
  pdf=render_pdf.build_pdf(ctx,'gl_budget_status.html.j2');(OUT/(key+'.pdf')).write_bytes(pdf)
  deck=render_fd.build_deck(ctx);(OUT/(key+'.pptx')).write_bytes(deck)
  doc=fitz.open(stream=pdf,filetype='pdf');slides=Presentation(BytesIO(deck))
  check(key+' PDF and PPT page counts match',len(doc)==len(slides.slides))
  text='\n'.join(p.get_text() for p in doc)
  check(key+' PDF retains all 30 department identifiers',all(f'D{i:02}' in text for i in range(30)))
  check(key+' all PPT slides contain report artwork',all(any(s.shape_type==13 for s in slide.shapes) for slide in slides.slides))
  check(key+' PDF contains no blank pages',all(len(p.get_text().strip())>50 for p in doc))
  doc[0].get_pixmap().save(OUT/(key+'.png'));doc.close()
(OUT/'results.json').write_text(json.dumps(results,indent=2))
raise SystemExit(0 if all(ok for _,ok in results) else 1)
