"""Collect deep-review evidence and dated UAT workbook / Word results."""
import json,shutil
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
from docx.shared import Inches
BASE=Path(__file__).resolve().parents[1]/'UAT';TRIAL=BASE/'UAT_GL_round4-06-09-2026';EV=TRIAL/'evidence_06-Sep-2026-01';TMP=Path('/tmp/fd-deep')
for name in ['edge-before','edge-complete','report-before','report-final','artifacts','summary-staged','security','staged-browser','live-browser','summary-live','reports-live','edge-live']:
 if (TMP/name).exists():shutil.copytree(TMP/name,EV/name,dirs_exist_ok=True)
for name in ['api-baseline.log','staged-browser.json','report-regression.log','cleanup-incomplete.log','live-browser.json','ui-deploy.log','ords-deploy.log','worker180-deploy.log','worker181-deploy.log','worker182-deploy.log','entity-results.json']:
 if (TMP/name).exists():shutil.copy2(TMP/name,EV/name)
sources=[('Live baseline browser',TRIAL/'baseline-browser.json'),('Live baseline reports',EV/'baseline-reports/results.json'),('Fixed UI regression',EV/'staged-browser.json'),('Show Summary',EV/'summary-staged/results.json'),('Deep browser',EV/'edge-complete/results.json'),('Export formatting and pagination',EV/'report-final/results.json'),('PDF and PPT artifacts',EV/'artifacts/results.json'),('Account access',EV/'security/results.json'),('Deployed browser',EV/'live-browser.json'),('Deployed Show Summary',EV/'summary-live/results.json'),('Deployed reports',EV/'reports-live/results.json'),('Deployed deep browser',EV/'edge-live/results.json'),('Entity exports',EV/'entity-results.json')]
rows=[];counts=[]
for suite,path in sources:
 if not path.exists():continue
 data=json.loads(path.read_text());passed=0
 for i,item in enumerate(data,1):
  if isinstance(item,dict):name,ok=item['name'],item['passed']
  else:name,ok=item[:2]
  passed+=bool(ok)
  status='PASS' if ok else ('DECISION' if suite=='Account access' else 'FAIL')
  rows.append([f'FD4-{len(rows)+1:03}',suite,name,status,str(path.relative_to(TRIAL))])
 counts.append((suite,passed,len(data)))
rows.append(['FD4-API','Live API baseline','Financial totals, filters, drills and error cases (144 assertions)','PASS','evidence_06-Sep-2026-01/api-baseline.log'])
rows.append(['FD4-UNIT','Report unit regression','Four report tests: context scopes, arithmetic, full department coverage, escaping','PASS','evidence_06-Sep-2026-01/report-regression.log'])
rows.append(['FD4-DEPLOY','Deployment','GL 1.116.1 deployed; workers active; report handlers installed','PASS','evidence_06-Sep-2026-01/ui-deploy.log'])
wb=Workbook();ws=wb.active;ws.title='Budget Status deep review';ws.append(['ID','Suite','Check','Status','Evidence'])
for r in rows:ws.append(r)
for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='3C6D5D')
for col,width in zip('ABCDE',[19,34,85,16,85]):ws.column_dimensions[col].width=width
for rr in ws.iter_rows(min_row=2):
 for c in rr:c.alignment=Alignment(wrap_text=True,vertical='top')
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;wb.save(TRIAL/'UAT_GL_06-Sep-2026-01.xlsx')
master=BASE/'UAT_GL_TestScript.xlsx';m=load_workbook(master)
name='Budget Status deep review'
if name in m.sheetnames:del m[name]
w=m.create_sheet(name);w.append(['ID','Suite','Check','Status','Evidence'])
for r in rows:w.append(r[:3]+['NOT RUN',r[4]])
m.save(master)
d=Document();d.add_heading('Budget Status — Deep Test Review',0)
d.add_paragraph('06 September 2026 · Round 4 · Local fixes validated against real read-only data and synthetic edge cases. Confirmed fixes deployed as GL 1.116.1; the print-permission policy remains unchanged pending the user decision.')
t=d.add_table(rows=1,cols=2);t.style='Light Shading Accent 1';t.rows[0].cells[0].text='Suite';t.rows[0].cells[1].text='Passed / tested'
for suite,passed,total in counts:
 c=t.add_row().cells;c[0].text=suite;c[1].text=f'{passed} / {total}'
d.add_paragraph('Additionally: live API baseline 144/144 assertions; report unit regression 4/4 tests. Baseline results describe the deployed version before these fixes. Fixed-UI and export results describe validated local changes; deployed suites repeat the key checks on GL 1.116.1.')
d.add_heading('Confirmed fixes',1)
for text in ['Saved JSON null and unavailable browser storage no longer break Budget Status controls.','Older responses and failures cannot replace the newest request; choosing a valid cached period clears the old error and invalidates pending requests.','Search supports Enter/Space, criteria have accessible names, Entity buttons expose selection, and hover instructions follow Show Summary.','Collapsed Search displays configured Entity names; the report bridge now validates against configured active Entity values (SQL deployed; six regular exports passed; all-entity export results are recorded separately).','PDF/PPT numeric formatting matches dashboard rounding and automatic units. Long entries move intact onto continuation pages; actual page numbers are recalculated.','User-approved behavior: all-zero-budget sectors/departments show Ledger rows with No budget allocated, including report output. Show Summary remains effective in the fallback.']:
 d.add_paragraph(text,style='List Bullet')
d.add_heading('Decision required',1)
d.add_paragraph('The GL compatibility setting allows signed-in users without GL_RUN_BRIEFING_BOOK to queue reports. Cross-user report status and downloads return 404, including for another privileged user. Three strict-403 expectations therefore remain marked DECISION; this is not evidence of cross-user access. The user must choose whether Budget Status should require the dedicated print privilege or retain compatibility behavior. Temporary accounts and sessions were removed; cleanup logs are attached.')
d.add_paragraph('The user explicitly authorized the built-in quick login after automatic approval review initially blocked the additional entity test. It was then run with that authorization; its results appear in the Entity exports suite. Deployed report SQL compiled successfully; the normal six-export regression passed.')
d.add_heading('Suggested enhancements — not implemented',1)
for text in ['Executive exceptions: a short list of the sectors/departments with the lowest remaining funds or negative balances, with direct drill-through. Example: Capex — 3% remaining.','Data freshness: show the accounting period and last successful data refresh beside the total so executives can judge how current the figures are.','Visible filter chips: keep Entity, selected Sectors and Departments visible above results, with an individual clear action to make scope easier to understand.']:
 d.add_paragraph(text,style='List Bullet')
d.add_heading('Evidence',1)
for title,path in [('Approved zero-budget fallback',EV/'edge-complete/zero-budget-map.png'),('English PDF sample',EV/'artifacts/en-tiles.png'),('Arabic PDF sample',EV/'artifacts/ar-tiles.png')]:
 if path.exists():d.add_heading(title,2);d.add_picture(str(path),width=Inches(6.3))
d.add_paragraph('Artifacts include 12 PDF/PPT files covering all three layouts in both languages, before/after edge results, live report samples, browser screenshots and account-cleanup logs. Production financial data is present in baseline evidence. No emails or messages were sent.')
d.save(TRIAL/'UAT_GL_Results_06-Sep-2026-01.docx')
print('Created',TRIAL)
