"""Live Budget Status report API verification; generates download-only test runs.

Uses the existing Admin quick-login credentials without logging secrets. No
recipient rules are created and no messages are sent. Save real export artifacts
and results in a new UAT round via FD_LIVE_EVIDENCE.
"""
import json,os,re,sys,time,urllib.request,urllib.error
from pathlib import Path
from io import BytesIO
import fitz
from pptx import Presentation

ORDS='https://gd5cec2eaeb21e3-prod.adb.me-abudhabi-1.oraclecloudapps.com/ords/admin'
ROOT=Path(__file__).resolve().parents[3]
OUT=Path(os.environ.get('FD_LIVE_EVIDENCE','/tmp/fd-live-report'));OUT.mkdir(parents=True,exist_ok=True)
results=[]
def check(name,ok,extra=''):
    results.append((name,bool(ok)));print(('PASS ' if ok else 'FAIL ')+name+' '+str(extra),flush=True)
def req(method,path,token=None,body=None,binary=False):
    headers={'Authorization':'Bearer '+token} if token else {}
    data=json.dumps(body).encode() if body is not None else None
    if data:headers['Content-Type']='application/json'
    r=urllib.request.Request(ORDS+path,data=data,method=method,headers=headers)
    try:
        with urllib.request.urlopen(r,timeout=120) as response:
            b=response.read();return response.status,b if binary else json.loads(b or b'{}')
    except urllib.error.HTTPError as e:
        b=e.read()
        try:return e.code,json.loads(b)
        except ValueError:return e.code,b.decode(errors='replace')[:200]

def main():
    auth=(ROOT/'final apps/Admin/Jet/js/services/authService.js').read_text()
    u,p=re.findall(r"username:\s*'([^']+)',\s*password:\s*'([^']+)'",auth)[0]
    status,session=req('POST','/dct/auth/login',body=dict(username=u,password=p))
    assert status==200 and session.get('sessionId'),'Login failed'
    token=session['sessionId']
    for path,method in [('/gl/fd/report','POST'),('/gl/fd/report/0','GET'),('/gl/fd/report/0/file','GET')]:
        s,_=req(method,path,body={} if method=='POST' else None);check(method+' '+path+' no token => 401',s==401,s)
    for body in [{},{'period':'13-2026'},{'period':'09-2026','format':'XLSX'},{'period':'09-2026','presentation':'bad'},{'period':'09-2026','entity':'bad'},{'period':'09-2026','sectors':"X' OR 1=1"}]:
        s,b=req('POST','/gl/fd/report',token,body);check('Invalid report parameter => 400 '+str(body),s==400,(s,b))
    for path in ['/gl/fd/report/0','/gl/fd/report/0/file']:
        s,_=req('GET',path,token);check('Unknown run => 404 '+path,s==404,s)
    _,filters=req('GET','/gl/actuals/filters',token);period=filters['defaultPeriod']
    s,cube=req('GET','/gl/fd/status?period='+period,token);assert s==200
    (OUT/'dashboard_cube.json').write_text(json.dumps(cube))
    candidates=[r for r in cube['rows'] if r['entity']=='DCT' and r['budget']>0]
    chosen=candidates[0]
    jobs=[]
    for layout in ('tiles','rows','map'):
        for fmt in ('PDF','PPTX'):
            params=dict(period=period,format=fmt,entity='ALL',presentation=layout,unit='X',lang='en',sector_sort='budget',department_sort='budget')
            if layout=='rows':params.update(entity='DCT',sectors=chosen['sector'],departments=chosen['costCenter'])
            if layout=='map':params.update(entity='MUSEUMS',lang='ar' if fmt=='PPTX' else 'en')
            s,b=req('POST','/gl/fd/report',token,params)
            check(layout+' '+fmt+' enqueued',s==200 and bool(b.get('runId')),(s,b))
            if s==200:jobs.append(dict(run_id=b['runId'],params=params,done=False))
    deadline=time.monotonic()+360
    while time.monotonic()<deadline and not all(j['done'] for j in jobs):
        for job in jobs:
            if job['done']:continue
            s,b=req('GET','/gl/fd/report/'+str(job['run_id']),token)
            if s!=200 or b.get('status') not in ('SUCCESS','FAILED'):continue
            job['done']=True;params=job['params'];key=params['presentation']+'_'+params['format']
            check(key+' report SUCCESS',b.get('status')=='SUCCESS',b)
            if b.get('status')!='SUCCESS':continue
            s,file=req('GET',f"/gl/fd/report/{job['run_id']}/file",token,binary=True)
            check(key+' authenticated file download',s==200 and isinstance(file,bytes),len(file))
            if s!=200:continue
            path=OUT/(key+('.pdf' if params['format']=='PDF' else '.pptx'));path.write_bytes(file)
            if params['format']=='PDF':
                doc=fitz.open(stream=file,filetype='pdf');text=doc[0].get_text()
                scoped=[r for r in cube['rows'] if (params['entity']=='ALL' or r['entity']==params['entity']) and (not params.get('sectors') or r['sector']==params['sectors']) and (not params.get('departments') or r['costCenter']==params['departments'])]
                for metric in ('budget','actual','encumbrance','fundsAvailable'):
                    expected=f"{sum(float(r.get(metric) or 0) for r in scoped):,.0f}"
                    check(key+' Total '+metric+' matches live dashboard',expected in text,expected)
                check(key+' title and selected period included','Budget Status' in text and period in text)
                check(key+' excludes Search/Print controls','Show Summary' not in text and 'Print' not in text)
                doc[0].get_pixmap().save(OUT/(key+'.png'));job['pages']=len(doc)
                doc.close()
            else:
                prs=Presentation(BytesIO(file));job['pages']=len(prs.slides)
                check(key+' deck opens with report slide images',len(prs.slides)>=3 and all(any(sh.shape_type==13 for sh in slide.shapes) for slide in prs.slides),len(prs.slides))
        if not all(j['done'] for j in jobs):time.sleep(3)
    check('All six report jobs finished',len(jobs)==6 and all(j['done'] for j in jobs))
    (OUT/'runs.json').write_text(json.dumps(jobs,indent=2));(OUT/'results.json').write_text(json.dumps(results,indent=2))
    print(f'{sum(ok for _,ok in results)}/{len(results)} passed',flush=True)
    return 0 if all(ok for _,ok in results) else 1
if __name__=='__main__':sys.exit(main())
