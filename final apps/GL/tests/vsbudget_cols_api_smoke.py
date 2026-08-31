#!/usr/bin/env python3
"""GL vs-Budget + Manage-columns API smoke (GL/db/35 + 33 re-run + 11 re-run,
reporting/db/25, runner _apply_sheet_cols — v1.96.0).

1. GET /gl/butil ships per-row budUtilPct / budVariance / budState (verdict of
   the displayed Total Actual vs the ADJUSTED ANNUAL budget — user rule
   2026-08-31: Annual always) + a budgetThresholds echo from the GL module
   settings BUD_UTIL_NEAR_PCT / BUD_UTIL_OVER_PCT; recomputed client-side the
   states must match exactly, under procash / costadj toggles too.
2. GET /gl/butil/agg carries the same verdict per Department / Sector group.
3. POST /gl/butil/xlsx forwards `sheetcols` (the Manage-columns saved view) to
   the runner's per-run sheet filter: the produced register sheet 1 holds
   EXACTLY the requested columns, in order, led by Budget Combination; a run
   WITHOUT sheetcols keeps the full sheet including the three new budget
   verdict columns.

Run: python vsbudget_cols_api_smoke.py     (defaults to the webtier)
"""
import io
import json
import os
import re
import ssl
import time
import urllib.request

import openpyxl

BASE = os.environ.get('GL_BASE', 'https://129.151.159.189')
AUTH = '/root/DCT-Task-Management/final apps/Admin/Jet/js/services/authService.js'
CTX = ssl.create_default_context()
CTX.check_hostname = False
CTX.verify_mode = ssl.CERT_NONE
YEAR = 2026

ok = fail = 0


def ck(name, cond, extra=''):
    global ok, fail
    print(('  PASS ' if cond else '  FAIL ') + name + ('  ' + str(extra) if extra else ''))
    if cond:
        ok += 1
    else:
        fail += 1


def login():
    u, p = re.findall(r"username:\s*'([^']+)',\s*password:\s*'([^']+)'", open(AUTH).read())[0]
    r = urllib.request.Request(BASE + '/ords/admin/dct/auth/login',
                               data=json.dumps({'username': u, 'password': p}).encode(),
                               method='POST', headers={'Content-Type': 'application/json'})
    return json.loads(urllib.request.urlopen(r, timeout=60, context=CTX).read())['sessionId']


TOK = login()


def get(path):
    rq = urllib.request.Request(BASE + '/ords/admin/gl' + path,
                                headers={'Authorization': 'Bearer ' + TOK})
    return json.loads(urllib.request.urlopen(rq, timeout=180, context=CTX).read())


def post(path, body):
    rq = urllib.request.Request(BASE + '/ords/admin/gl' + path,
                                data=json.dumps(body).encode(), method='POST',
                                headers={'Authorization': 'Bearer ' + TOK,
                                         'Content-Type': 'application/json'})
    return json.loads(urllib.request.urlopen(rq, timeout=120, context=CTX).read())


def verdict(act, ba, th):
    if abs(ba) > 0.005:
        pct = 100.0 * act / ba
        return 'OVER' if pct > th['over'] else 'NEAR' if pct >= th['near'] else 'OK'
    return 'OVER' if act > 0.005 else 'NOBUDGET'


def parity(name, query, procash=False):
    d = get('/butil?' + query)
    th = d.get('budgetThresholds') or {}
    bad = miss = 0
    for r in d['items']:
        st = r.get('budState')
        if st is None:
            miss += 1
            continue
        act = (r.get('actualAp') or 0) + (r.get('actualGrn') or 0) \
            + ((r.get('procash') or 0) if procash else 0)
        ba = r.get('budgetAnnual') or 0
        exp = verdict(act, ba, th)
        if exp != st:
            bad += 1
        elif st != 'NOBUDGET' and abs(ba) > 0.005:
            if abs(round(100.0 * act / ba, 1) - (r.get('budUtilPct') or 0)) > 0.06:
                bad += 1
            elif abs(round(act - ba, 2) - (r.get('budVariance') or 0)) > 0.05:
                bad += 1
    ck(name, bad == 0 and miss == 0 and len(d['items']) > 0,
       '%d rows, %d bad, %d missing' % (len(d['items']), bad, miss))
    return d


print('== thresholds echo ==')
d = parity('per-row parity (default toggles)', 'year=%d&limit=500' % YEAR)
th = d.get('budgetThresholds') or {}
ck('budgetThresholds echoed', th.get('near') is not None and th.get('over') is not None, th)
states = {}
for r in d['items']:
    states[r.get('budState')] = states.get(r.get('budState'), 0) + 1
ck('every row classified', None not in states, states)

print('== toggle parity ==')
parity('per-row parity (procash=Y)', 'year=%d&limit=300&procash=Y' % YEAR, procash=True)
parity('per-row parity (costadj=N)', 'year=%d&limit=300&costadj=N' % YEAR)
parity('per-row parity (period YTD)', 'year=%d&limit=300&period=06-%d' % (YEAR, YEAR))

print('== agg tabs ==')
for lvl in ('dept', 'sector'):
    a = get('/butil/agg?level=%s&year=%d' % (lvl, YEAR))
    tha = a.get('budgetThresholds') or {}
    bad = 0
    for g in a['items']:
        st = g.get('budState')
        act = (g.get('actualAp') or 0) + (g.get('actualGrn') or 0)
        if st is None or st != verdict(act, g.get('budgetAnnual') or 0, tha):
            bad += 1
    ck('agg %s parity' % lvl, bad == 0 and len(a['items']) > 0,
       '%d groups, %d bad' % (len(a['items']), bad))

print('== register sheetcols (runner _apply_sheet_cols) ==')


def run_register(body):
    rid = post('/butil/xlsx', body)['runId']
    for _ in range(70):
        time.sleep(5)
        s = get('/butil/xlsx/%d' % rid)
        if s['status'] == 'SUCCESS' and s.get('hasFile'):
            rq = urllib.request.Request(BASE + '/ords/admin/gl/butil/xlsx/%d/file' % rid,
                                        headers={'Authorization': 'Bearer ' + TOK})
            return openpyxl.load_workbook(
                io.BytesIO(urllib.request.urlopen(rq, timeout=300, context=CTX).read()))
        if s['status'] == 'FAILED':
            raise RuntimeError('run %d FAILED: %s' % (rid, s.get('error')))
    raise RuntimeError('run %d timed out' % rid)


def headers(ws):
    out = []
    for c in ws[4]:
        if c.value in (None, ''):
            break
        out.append(str(c.value))
    return out


SCOPE = {'year': YEAR, 'project': '4511000015', 'cmtmode': 'NONE'}

want = 'budget_combination,project_number,project_name,expenditure_type,annual_budget,' \
       'budget_utilization_pct,budget_variance,budget_status,actual_ap,actual_grn,fund_available'
wb = run_register(dict(SCOPE, sheetcols=want))
hs = headers(wb.worksheets[0])
exp = [w.replace('_', ' ').title() for w in want.split(',')]
ck('filtered sheet-1 columns exact + ordered', hs == exp, hs)
stat_i = exp.index('Budget Status')
vals = set()
for row in wb.worksheets[0].iter_rows(min_row=5, values_only=True):
    if row[0] is None:
        break
    vals.add(str(row[stat_i]).split(' ', 1)[-1] if row[stat_i] else '')
ck('budget_status values render', len(vals) > 0, vals)
ck('other sheets untouched', len(wb.worksheets) > 1, len(wb.worksheets))

wb2 = run_register(dict(SCOPE))
hs2 = headers(wb2.worksheets[0])
ck('full sheet keeps every column', 'Budget Utilization Pct' in hs2 and 'Budget Variance' in hs2
   and 'Budget Status' in hs2 and 'Plan Status' in hs2 and 'Budget Combination' in hs2,
   '%d columns' % len(hs2))
ck('new budget columns after Ytd Budget', hs2.index('Budget Utilization Pct') == hs2.index('Ytd Budget') + 1)

print('\n%d/%d passed' % (ok, ok + fail))
raise SystemExit(0 if fail == 0 else 1)
