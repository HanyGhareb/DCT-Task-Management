"""Budget Status: shared paginated dashboard artwork for PDF and PowerPoint.

The query uses the live dashboard's GL chapter/budget-group pairs. Aggregate
in Python using the same contextual scopes as the screen: sectors = entity;
departments = entity + sectors; chapter bands = entity + sectors + departments.
No top-N truncation. PowerPoint slides embed the same rendered page artwork.
"""
from collections import defaultdict
from io import BytesIO
from decimal import Decimal, ROUND_HALF_UP
import math

MEASURES = ('budget', 'actual', 'encumbrance', 'funds_available')
EN = dict(title='Budget Status', chapters='Chapters', sectors='Sectors', departments='Departments (Cost Centres)',
          total='Total', budget='Budget', actual='Actual', encumbrance='Encumbrance', funds_available='Fund available',
          all='All', tiles='Composition tiles', rows='Ledger rows', map='Proportional map',
          ok='On track', tight='Tight', over='Over committed', no_data='No data for these filters',
          legend='Area = budget · Colour = consumed share (Actual + Encumbrance). Full figures follow.',
          selected='Selected', search='Department search', unit='Figures in',
          context='Sector figures follow Entity; department figures also follow selected Sectors; chapter totals follow all selections.',
          zero='Unbudgeted activity receives a small minimum area; all entries are listed on the following pages.')
AR = dict(title='حالة الموازنة', chapters='الأبواب', sectors='القطاعات', departments='الإدارات (مراكز التكلفة)',
          total='الإجمالي', budget='الموازنة', actual='الفعلي', encumbrance='الارتباطات', funds_available='الرصيد المتاح',
          all='الكل', tiles='بطاقات التكوين', rows='صفوف دفترية', map='خريطة نسبية',
          ok='ضمن المسار', tight='محدود', over='تجاوز الارتباط', no_data='لا توجد بيانات لهذه المعايير',
          legend='المساحة = الموازنة · اللون = نسبة المستخدم (الفعلي + الارتباطات). التفاصيل في الصفحات التالية.',
          selected='المحدد', search='بحث الإدارات', unit='عرض الأرقام',
          context='القطاعات حسب الجهة؛ الإدارات حسب القطاعات المحددة أيضاً؛ إجماليات الأبواب حسب جميع الاختيارات.',
          zero='تمنح الأنشطة بلا موازنة مساحة دنيا صغيرة؛ جميع البنود مدرجة في الصفحات التالية.')
CH_EN = ['Payroll', 'Opex', 'Capex', 'Subsidy', 'Aids & Grants']
CH_AR = ['الرواتب', 'التشغيلية', 'الرأسمالية', 'الإعانات', 'المساعدات والمنح']


def _sum(rows):
    return {k: sum(float(r.get(k) or 0) for r in rows) for k in MEASURES}


def _group(rows, key, name, ar=False):
    groups = defaultdict(list)
    for r in rows:
        groups[r[key]].append(r)
    return [dict(_sum(rr), code=code,
                 name=(rr[0].get('sector_name_ar') if ar and key == 'sector' else None) or rr[0].get(name) or code)
            for code, rr in groups.items() if any(_sum(rr).values())]


def _pct(value, budget):
    return math.floor(value / budget * 100 + .5) if budget else 0


def _sort(rows, key):
    if key == 'name':
        return sorted(rows, key=lambda r: r['name'])
    metric = {'used': 'actual', 'free': 'funds_available'}.get(key)
    return sorted(rows, key=lambda r: (_pct(r[metric], r['budget']), r['budget']) if metric else r['budget'], reverse=True)


def _rects(items, x=0, y=0, w=1184, h=420):
    """Balanced binary treemap; positive budgets preserve exact area ratios."""
    if not items:
        return []
    if len(items) == 1:
        return [dict(items[0], x=x, y=y, w=w, h=h)]
    total = sum(r.get('weight', r['budget']) for r in items)
    acc, cut = 0, 1
    for i, r in enumerate(items[:-1], 1):
        acc += r.get('weight', r['budget'])
        cut = i
        if acc >= total / 2:
            break
    ratio = acc / total
    if w >= h:
        return _rects(items[:cut], x,y,w*ratio,h) + _rects(items[cut:],x+w*ratio,y,w*(1-ratio),h)
    return _rects(items[:cut],x,y,w,h*ratio) + _rects(items[cut:],x,y+h*ratio,w,h*(1-ratio))


def prepare(ctx):
    params = ctx.get('params') or {}
    ar = params.get('lang') == 'ar'
    labels = AR if ar else EN
    sections = {s['key']: s.get('data') or [dict(zip(s['columns'], r)) for r in s.get('rows', [])]
                for s in ctx.get('sections') or []}
    cube = sections.get('cube', [])
    sectors = set(filter(None, (params.get('sectors') or '').split('|')))
    depts = set(filter(None, (params.get('departments') or '').split('|')))
    dept_cube = [r for r in cube if not sectors or r['sector'] in sectors]
    band_cube = [r for r in dept_cube if not depts or r['department'] in depts]
    sec = _group(cube, 'sector', 'sector_name', ar)
    dep = _group(dept_cube, 'department', 'department_name', ar)
    q = (params.get('department_search') or '').strip().lower()
    dep = [r for r in dep if not q or q in (r['code'] + ' ' + r['name']).lower()]
    thresholds = {r['setting_key']: r['setting_value'] for r in sections.get('thresholds', [])}
    def setting(key, fallback):
        try:
            v = float(thresholds.get(key, fallback))
            return v if math.isfinite(v) else fallback
        except (TypeError, ValueError):
            return fallback
    near, over = setting('BUD_UTIL_NEAR_PCT',90), setting('BUD_UTIL_OVER_PCT',100)
    unit = params.get('unit', 'auto')
    def number(v):
        v = float(v or 0)
        sign = '-' if v < 0 else ''
        magnitude = abs(v)
        if unit == 'X':
            rounded = math.floor(v + .5)
            return ('-' if v < 0 and rounded == 0 else '') + f'{rounded:,}'
        def fixed(value, places, grouped=False):
            # Intl uses the shortest decimal value; toFixed rounds the binary value.
            d = Decimal(str(value)) if grouped else Decimal.from_float(value)
            d = d.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
            return format(d, (',' if grouped else '') + '.' + str(places) + 'f')
        if unit == 'B': return sign + fixed(magnitude/1e9, 2, True) + 'B'
        if unit == 'M': return sign + fixed(magnitude/1e6, 1, True) + 'M'
        if unit == 'K': return sign + f'{math.floor(magnitude/1e3 + .5):,}K'
        if magnitude >= 1e9: return sign + fixed(magnitude/1e9, 2) + 'B'
        if magnitude >= 1e6: return sign + fixed(magnitude/1e6, 1) + 'M'
        if magnitude >= 1e3: return sign + f'{math.floor(magnitude/1e3 + .5)}K'
        return sign + str(math.floor(magnitude + .5))
    def decorate(r):
        r = dict(r)
        b = r['budget']; consumed = (r['actual']+r['encumbrance'])/b*100 if b else 0
        state = 'over' if r['funds_available'] < 0 or consumed > over+.0001 else 'tight' if consumed >= near else 'ok'
        heat = 4 if state=='over' else 3 if state=='tight' else 2 if consumed>=max(60,near-10) else 1 if consumed>=60 else 0
        r.update(state=state,state_label=labels[state],consumed=consumed,dark=heat>=2,
                 color=['#DCE9E3','#9CC3B3','#4F8A75','#2C5044','#9E3B33'][heat])
        r['figures'] = [dict(key=k, label=labels[k], text=number(r[k]),
                            pct=_pct(r[k],b), dash=max(0,min(100,r[k]/b*100)) if b>0 else 0) for k in MEASURES]
        a = max(0,min(100,r['actual']/b*100)) if b>0 else 0
        e = max(0,min(100-a,r['encumbrance']/b*100)) if b>0 else 0
        f = max(0,min(100-a-e,r['funds_available']/b*100)) if b>0 else 0
        r['segments']=[a,e,f]
        return r
    bands = [decorate(dict(_sum(band_cube),name=labels['total'],code='TOTAL'))]
    for i in range(1,6):
        rows = [r for r in band_cube if r['chapter']==f'CH{i}']
        name = (f'الباب {i} · {CH_AR[i-1]}' if ar else f'Chapter {i} · {CH_EN[i-1]}')
        bands.append(decorate(dict(_sum(rows),name=name,code=f'CH{i}')))
    pages=[]
    def chunks(title, kind, rows, count):
        for i in range(0,len(rows),count):
            pages.append(dict(title=title,kind=kind,rows=rows[i:i+count]))
    chunks(labels['chapters'],'bands',bands,3)
    presentation=params.get('presentation','tiles')
    for key, rows, selection in [('sectors',sec,sectors),('departments',dep,depts)]:
        sort_key=params.get('sector_sort' if key=='sectors' else 'department_sort','budget')
        rows=[decorate(r) for r in _sort(rows, 'budget' if presentation=='map' else sort_key)]
        max_b=max([r['budget'] for r in rows]+[1])
        for r in rows:
            r['selected']=r['code'] in selection
            r['scale']=max(0,r['budget'])/max_b*100
        if not rows:
            pages.append(dict(title=labels[key],kind='empty',rows=[]))
        elif presentation=='map' and all(not r['budget'] for r in rows):
            chunks(labels[key] + ' · ' + ('لا توجد موازنة مخصصة' if ar else 'No budget allocated'), 'rows', rows, 8)
        elif presentation=='map':
            floor=max_b*.006
            weighted=[dict(r,weight=r['budget'] if r['budget']>0 else floor) for r in rows]
            pages.append(dict(title=labels[key],kind='map',rows=_rects(weighted)))
            chunks(labels[key], 'rows', rows, 8) # full names/values even for the tiniest map cell
        else:
            chunks(labels[key],presentation,rows,9 if presentation=='tiles' else 8)
    excluded=sections.get('excluded', [])
    if excluded:
        basis = ('الأساس: أرصدة الأستاذ حتى الفترة المحددة. الأبواب 1–3 / مجموعة الموازنة 1؛ الباب 4 / المجموعة 3؛ الباب 5 / المجموعة 5.' if ar else
                 'Basis: GL balances, year to date at the selected period. Chapters 1–3 / Budget Group 1; Chapter 4 / Group 3; Chapter 5 / Group 5.')
        explanation = ('الأرصدة المستبعدة من مجموعة الموازنة 1 خارج الأبواب المشمولة (جميع الجهات؛ ليست جزءاً من إجمالي التحديد):' if ar else
                       'Excluded Budget Group 1 balances outside the included chapters (all entities; not part of the selected total):')
        note_rows=[basis, explanation]+[labels[k] + ': ' + number(excluded[0].get(k)) for k in MEASURES]
        pages.append(dict(title='نطاق التقرير' if ar else 'Report scope',kind='scope',rows=note_rows))
    entity=params.get('entity') or 'ALL'
    ent={'ALL':labels['all'],'MUSEUMS':'MSS','MASTERPIECES': 'Masterpieces'}.get(entity,entity)
    unit_name={'X':'Exact number' if not ar else 'الرقم الكامل','auto':'Automatic' if not ar else 'تلقائي',
               'K':'Thousands' if not ar else 'آلاف','M':'Millions' if not ar else 'ملايين','B':'Billions' if not ar else 'مليارات'}.get(unit,unit)
    scope=' · '.join([str(params.get('period') or ctx.get('period') or ''), ent, labels.get(presentation,presentation), unit_name])
    sel_names=[r['name'] for r in sec if r['code'] in sectors]
    # Department search only filters the cards, never the selection summary or bands.
    dep_names=[r['name'] for r in _group(dept_cube,'department','department_name',ar) if r['code'] in depts]
    selections=f"{labels['sectors']}: {', '.join(sel_names) or labels['all']} · {labels['departments']}: {', '.join(dep_names) or labels['all']}"
    if q: selections+=f" · {labels['search']}: {params['department_search']}"
    # Long selections get their own continuation pages rather than clipping report headers.
    if len(selections)>260:
        parts=[selections[i:i+180] for i in range(0,len(selections),180)]
        for i in range(0,len(parts),8):
            pages.append(dict(title=labels['selected'],kind='scope',rows=parts[i:i+8]))
        selections=f"{labels['selected']}: {len(sectors)} {labels['sectors']} · {len(depts)} {labels['departments']}"
    return dict(ctx, fd=dict(pages=pages,labels=labels,scope=scope,selections=selections,ar=ar,near=near,over=over))


def build_deck(ctx):
    from playwright.sync_api import sync_playwright
    from pptx import Presentation
    from pptx.util import Inches
    import render_pdf
    html=render_pdf.render_html(ctx,'gl_budget_status.html.j2')
    prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
    with sync_playwright() as p:
        browser=p.chromium.launch(args=['--no-sandbox'])
        try:
            page=browser.new_page(viewport={'width':1280,'height':720},device_scale_factor=1.5)
            page.set_content(html,wait_until='networkidle')
            for panel in page.locator('.report-page').all():
                slide=prs.slides.add_slide(prs.slide_layouts[6])
                slide.shapes.add_picture(BytesIO(panel.screenshot()),0,0,width=prs.slide_width,height=prs.slide_height)
        finally:
            browser.close()
    out=BytesIO(); prs.save(out); return out.getvalue()
